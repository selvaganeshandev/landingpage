"""
Backlink profile — .xlsx export.

A Summary sheet carrying the whole profile at a glance, then one sheet per
stored list. The summary comes first deliberately: the link table alone cannot
answer "how big is this profile", and the two numbers genuinely differ (see
`Backlinks listed individually` below), so the spreadsheet has to say which is
which rather than leave the reader to infer it.
"""
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Matches the palette used by the opportunities export and the report export.
HEADER_FILL = PatternFill(start_color="FFD966", fill_type="solid")
TITLE_FILL = PatternFill(start_color="00B050", fill_type="solid")
HEADER_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, color="FFFFFF", size=12)

# Every breakdown DataForSEO returns pre-aggregated on the summary call. The
# blank key means "unattributed" and is relabelled rather than dropped, so the
# figures still add up to the profile total.
BREAKDOWNS = [
    ('referring_links_tld', 'By source TLD'),
    ('referring_links_types', 'By link type'),
    ('referring_links_attributes', 'By link attribute'),
    ('referring_links_platform_types', 'By platform type'),
    ('referring_links_semantic_locations', 'By placement on page'),
    ('referring_links_countries', 'By country'),
]

BACKLINK_COLUMNS = [
    ('Source domain', lambda r: r.domain_from),
    ('Source URL', lambda r: r.url_from),
    ('Source page title', lambda r: r.page_from_title),
    ('Target URL', lambda r: r.url_to),
    ('Anchor', lambda r: r.anchor),
    ('Link type', lambda r: r.item_type),
    ('Follow', lambda r: 'dofollow' if r.dofollow else 'nofollow'),
    ('Domain rank', lambda r: r.domain_from_rank),
    ('Page rank', lambda r: r.page_from_rank),
    ('Link rank', lambda r: r.rank),
    ('Spam score', lambda r: r.backlink_spam_score),
    ('Country', lambda r: r.domain_from_country),
    ('Placement', lambda r: r.semantic_location),
    ('New', lambda r: 'yes' if r.is_new else ''),
    ('Lost', lambda r: 'yes' if r.is_lost else ''),
    ('Broken', lambda r: 'yes' if r.is_broken else ''),
    ('First seen', lambda r: r.first_seen.date() if r.first_seen else ''),
    ('Last seen', lambda r: r.last_seen.date() if r.last_seen else ''),
]

REFERRING_DOMAIN_COLUMNS = [
    ('Domain', lambda r: r.domain_name),
    ('Domain rank', lambda r: r.rank),
    ('Backlinks', lambda r: r.backlinks),
    ('Referring pages', lambda r: r.referring_pages),
    ('Broken backlinks', lambda r: r.broken_backlinks),
    ('Spam score', lambda r: r.backlinks_spam_score),
    ('Country', lambda r: r.country),
    ('New', lambda r: 'yes' if r.is_new else ''),
    ('Lost', lambda r: 'yes' if r.is_lost else ''),
    ('First seen', lambda r: r.first_seen.date() if r.first_seen else ''),
]

ANCHOR_COLUMNS = [
    ('Anchor text', lambda r: r.anchor),
    ('Backlinks', lambda r: r.backlinks),
    ('Referring domains', lambda r: r.referring_domains),
    ('Referring pages', lambda r: r.referring_pages),
    ('Spam score', lambda r: r.backlinks_spam_score),
    ('First seen', lambda r: r.first_seen.date() if r.first_seen else ''),
]

PAGE_COLUMNS = [
    ('Page', lambda r: r.page_url),
    ('Backlinks', lambda r: r.backlinks),
    ('Referring domains', lambda r: r.referring_domains),
    ('Referring pages', lambda r: r.referring_pages),
    ('Status code', lambda r: r.status_code or ''),
    ('First seen', lambda r: r.first_seen.date() if r.first_seen else ''),
]


def _title(ws, text, span):
    ws.cell(row=1, column=1, value=text).font = TITLE_FONT
    ws.cell(row=1, column=1).fill = TITLE_FILL
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)


def _write_table(wb, title, columns, rows, subtitle=''):
    ws = wb.create_sheet(title=title)
    _title(ws, title, len(columns))
    if subtitle:
        ws.cell(row=2, column=1, value=subtitle)

    for col, (heading, _) in enumerate(columns, start=1):
        cell = ws.cell(row=3, column=col, value=heading)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)

    for i, row in enumerate(rows, start=4):
        for col, (_, getter) in enumerate(columns, start=1):
            ws.cell(row=i, column=col, value=getter(row))

    # Width from the widest cell, clamped so a long URL column cannot push the
    # rest off screen.
    for col, (heading, getter) in enumerate(columns, start=1):
        longest = len(heading)
        for row in rows:
            longest = max(longest, len(str(getter(row))))
        ws.column_dimensions[get_column_letter(col)].width = min(max(longest + 2, 10), 60)

    ws.freeze_panes = 'A4'
    return ws


def _write_summary(wb, snapshot, domain_label):
    ws = wb.create_sheet(title='Summary')
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 62

    _title(ws, f'Backlink profile — {domain_label}', 3)

    row = 3

    def section(label):
        nonlocal row
        row += 1
        cell = ws.cell(row=row, column=1, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        ws.cell(row=row, column=2).fill = HEADER_FILL
        ws.cell(row=row, column=3).fill = HEADER_FILL
        row += 1

    def line(label, value, note=''):
        nonlocal row
        ws.cell(row=row, column=1, value=label).font = HEADER_FONT
        ws.cell(row=row, column=2, value=value)
        if note:
            ws.cell(row=row, column=3, value=note).alignment = Alignment(wrap_text=True)
        row += 1

    stored = snapshot.items.count()

    section('Profile')
    line('Domain', domain_label)
    line('Fetched', snapshot.completed_at.strftime('%d %b %Y %H:%M') if snapshot.completed_at else '')
    line('First link seen', snapshot.first_seen.strftime('%d %b %Y') if snapshot.first_seen else '')
    line('Domain rank', snapshot.rank, 'DataForSEO authority score, 0-1000. Higher is stronger.')
    line('Spam score', f'{snapshot.backlinks_spam_score}%',
         'Share of the profile coming from low-quality sources, 0-100.')

    section('Volume')
    line('Backlinks', snapshot.backlinks, 'Every inbound link found, including many from one site.')
    line('Backlinks listed individually', stored,
         'How many appear on the Backlinks sheet. DataForSEO counts links in the total that it '
         'will not itemise, so this is normally the smaller number — it is not a truncation of '
         'your data unless the row below says otherwise.')
    line('Detail capped', 'yes' if snapshot.is_truncated else 'no',
         'Yes means the profile was larger than the per-fetch storage cap and only the '
         'highest-ranked links were kept. The totals on this sheet are unaffected.')
    line('Referring domains', snapshot.referring_main_domains,
         'Distinct websites linking to you — usually a better measure of reach than raw backlinks.')
    line('Referring domains (nofollow)', snapshot.referring_main_domains_nofollow)
    line('Referring pages', snapshot.referring_pages)
    line('Referring IPs', snapshot.referring_ips)
    line('Referring subnets', snapshot.referring_subnets,
         'Many links from few subnets often means one network rather than independent coverage.')

    section('Problems')
    line('Broken backlinks', snapshot.broken_backlinks,
         'Links pointing at a page that no longer loads — earned authority being thrown away.')
    line('Broken pages', snapshot.broken_pages)

    section('Your site')
    line('Crawled pages', snapshot.crawled_pages)
    line('Internal links', snapshot.internal_links_count)
    line('External links', snapshot.external_links_count)

    for field, label in BREAKDOWNS:
        data = getattr(snapshot, field, None) or {}
        if not data:
            continue
        section(label)
        for key, value in sorted(data.items(), key=lambda kv: kv[1], reverse=True):
            ws.cell(row=row, column=1, value=key or '(unattributed)')
            ws.cell(row=row, column=2, value=value)
            row += 1

    return ws


def build_backlinks_workbook(snapshot, domain_label) -> BytesIO:
    """Workbook for one completed snapshot. Caller owns the HTTP response."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    _write_summary(wb, snapshot, domain_label)

    items = list(snapshot.items.all().order_by('-rank'))
    _write_table(
        wb, 'Backlinks', BACKLINK_COLUMNS, items,
        subtitle=(
            f'{len(items)} links listed individually of {snapshot.backlinks} counted in the '
            f'profile — see the Summary sheet.'
        ),
    )
    _write_table(
        wb, 'Referring Domains', REFERRING_DOMAIN_COLUMNS,
        list(snapshot.referring_domain_rows.all().order_by('-rank')),
    )
    _write_table(
        wb, 'Anchors', ANCHOR_COLUMNS,
        list(snapshot.anchor_rows.all().order_by('-backlinks')),
    )
    _write_table(
        wb, 'Top Pages', PAGE_COLUMNS,
        list(snapshot.page_rows.all().order_by('-backlinks')),
        subtitle='Your own pages, ranked by how many backlinks they receive.',
    )

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
