"""
Citations export — the whole Citations page as a multi-sheet workbook.

The on-screen table is paginated ten rows at a time, so a domain with a thousand
citations is effectively unreadable there. This writes every citation out once,
plus the aggregates the page shows above the table, so the file stands on its own
without needing the app open beside it.

Sheet layout mirrors the page top-to-bottom:

  Summary     the eight metric cards, the status breakdown, and what each means
  Citations   one row per citation event — the full list, unpaginated
  By Source   one row per cited hostname, ranked
  By Platform citations attributed to each AI platform

Deliberately one row per *citation event*, not per unique URL: a URL cited in
five answers is five rows. That is what "Total Citations" counts on the page, and
an export that silently deduplicated would not reconcile with it. Use the By
Source sheet for the deduplicated view.
"""
import logging
from datetime import datetime
from urllib.parse import urlparse

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from core.queryset_scoping import user_can_access_domain
from domains.models import Domain
from prompts.models import PromptAnalytics

from .models import CitationURL

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=13)
BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'),
)

# Same limit the mentions export uses: Excel rejects any cell over 32,767 chars
# and openpyxl raises rather than truncating, failing the whole workbook.
CELL_LIMIT = 32000


def _cell_text(value, limit=CELL_LIMIT):
    text = '' if value is None else str(value)
    return text if len(text) <= limit else text[:limit] + '… [truncated]'


def _naive(dt):
    """Excel cannot store a timezone-aware datetime; write a real date so the
    column sorts and filters as a date rather than as a string."""
    if not dt:
        return ''
    return timezone.localtime(dt).replace(tzinfo=None) if timezone.is_aware(dt) else dt


def _strip_www(host):
    """Drop a leading 'www.' — as a prefix, not as a character set.

    `lstrip('www.')` would eat any leading run of 'w' and '.', turning
    'wwf.org' into 'f.org'.
    """
    return host[4:] if host.startswith('www.') else host


def _host(url):
    try:
        return _strip_www((urlparse(url).netloc or '').lower())
    except Exception:
        return ''


def _write_headers(ws, headers, row=1):
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER


def _autosize(ws, max_width=70):
    """Width from the longest cell, capped. Uncapped, a full URL column pushes
    every later column off the visible page."""
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        longest = 0
        for cell in column_cells:
            if cell.value is not None:
                longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), max_width)


def _finish_table(ws, n_rows, n_cols, header_row=1):
    """Freeze the header and enable filtering — a 1,200-row sheet is unusable
    without both."""
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate
    if n_rows:
        ws.auto_filter.ref = f'A{header_row}:{get_column_letter(n_cols)}{header_row + n_rows}'
    _autosize(ws)


def _status_for(crawled):
    """Map a crawl row to the status shown on the page.

    Mirrors citations_list exactly, including the ordering quirk: 'blocked' rows
    carry http_status_code=403, so blocked must be tested before the generic
    >=400 branch or they would all read as broken.
    """
    if not crawled:
        return 'Pending', None
    crawl_status = crawled.crawl_status
    code = crawled.http_status_code
    if crawl_status == 'success' and code and code < 400:
        return 'Valid', code
    if crawl_status == 'blocked':
        return 'Blocked', code
    if (code and code >= 400) or crawl_status in ('failed', 'broken'):
        return 'Broken', code
    return 'Pending', code


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def citations_export(request):
    """
    Export every citation for a domain as a multi-sheet .xlsx.

    Query params:
        domain_id: Required
    """
    domain_id = request.query_params.get('domain_id')
    if not domain_id:
        return Response({'error': 'domain_id is required'}, status=status.HTTP_400_BAD_REQUEST)

    # 404 rather than 403 so the response cannot be used to confirm a domain exists.
    if not user_can_access_domain(request.user, domain_id, request):
        return Response({'error': 'Domain not found'}, status=status.HTTP_404_NOT_FOUND)

    try:
        domain = Domain.objects.get(id=domain_id)
    except Domain.DoesNotExist:
        return Response({'error': 'Domain not found'}, status=status.HTTP_404_NOT_FOUND)

    domain_host = (domain.url or '').replace('https://', '').replace('http://', '').rstrip('/')

    # One pass over the analytics, one over the crawl rows. Everything else is
    # computed in memory — the page's own endpoint re-queries per page, which is
    # affordable at ten rows and is not at a thousand.
    analytics = (
        PromptAnalytics.objects
        .filter(prompt__group__domain=domain, track_status='COMP')
        .exclude(citation_list=[])
        .exclude(citation_list__isnull=True)
        .select_related('prompt', 'prompt__group')
        .order_by('-created_at')
    )

    crawl_map = {}
    for row in CitationURL.objects.filter(domain=domain).only(
        'url', 'crawl_status', 'http_status_code', 'last_crawled_at'
    ):
        crawl_map[row.url] = row
        crawl_map[row.url.rstrip('/')] = row

    rows = []
    by_source = {}
    by_platform = {}
    status_counts = {'Valid': 0, 'Broken': 0, 'Blocked': 0, 'Pending': 0}

    for pa in analytics:
        citations = pa.citation_list or []
        if not isinstance(citations, list):
            continue
        prompt_text = getattr(pa.prompt, 'prompt', '') or ''
        # PromptGroup has no `name` — groups are identified by `theme` here and
        # everywhere else in the UI.
        group_name = getattr(getattr(pa.prompt, 'group', None), 'theme', '') or ''

        for position, url in enumerate(citations, start=1):
            if not url:
                continue
            crawled = crawl_map.get(url) or crawl_map.get(str(url).rstrip('/'))
            label, code = _status_for(crawled)
            status_counts[label] = status_counts.get(label, 0) + 1

            host = _host(url)
            # Same substring test citations_list uses, so the Type column
            # reconciles with the page's Your Domain / Third Party tabs.
            is_own = bool(domain_host) and domain_host in (url or '')

            rows.append([
                len(rows) + 1,
                label,
                code if code else '',
                _cell_text(url),
                host,
                'Your Domain' if is_own else 'Third Party',
                (pa.platform or '').title(),
                position,
                _cell_text(prompt_text),
                _cell_text(group_name),
                _naive(pa.created_at),
                _naive(getattr(crawled, 'last_crawled_at', None)),
            ])

            bucket = by_source.setdefault(host, {'count': 0, 'own': is_own, 'platforms': set(), 'last': None})
            bucket['count'] += 1
            if pa.platform:
                bucket['platforms'].add((pa.platform or '').title())
            if bucket['last'] is None or (pa.created_at and pa.created_at > bucket['last']):
                bucket['last'] = pa.created_at

            platform_label = (pa.platform or 'Unknown').title()
            by_platform[platform_label] = by_platform.get(platform_label, 0) + 1

    total = len(rows)
    own_total = sum(1 for r in rows if r[5] == 'Your Domain')
    response_count = analytics.count()
    all_completed = PromptAnalytics.objects.filter(
        prompt__group__domain=domain, track_status='COMP'
    ).count()

    wb = Workbook()

    # ---------- Sheet 1: Summary ----------
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = f'Citations Export — {domain.name}'
    ws['A1'].font = TITLE_FONT
    ws['A2'] = f'Generated {timezone.localtime(timezone.now()).strftime("%d %b %Y, %H:%M")}'
    ws['A3'] = f'Domain: {domain.url or ""}'

    _write_headers(ws, ['Metric', 'Value', 'What it means'], row=5)
    citation_rate = round(response_count / all_completed * 100, 1) if all_completed else 0
    avg_per_response = round(total / all_completed, 1) if all_completed else 0
    summary_rows = [
        ('Total Citations', total,
         'Every citation event. A URL cited in five answers counts five times.'),
        ('Unique Sources', len(by_source),
         'Distinct hostnames cited.'),
        ('Your Domain', own_total,
         'Citations pointing at your own website.'),
        ('Third Party', total - own_total,
         'Citations pointing anywhere else.'),
        ('Citation Rate', f'{citation_rate}%',
         'Share of completed responses carrying at least one citation.'),
        ('Avg Citations/Response', avg_per_response,
         'Total citations divided by all completed responses, including those with none.'),
        ('Responses With Citations', response_count, 'Completed responses that cited at least one source.'),
        ('Completed Responses', all_completed, 'All completed responses for this domain.'),
    ]
    for i, (metric, value, meaning) in enumerate(summary_rows, start=6):
        ws.cell(row=i, column=1, value=metric).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)
        ws.cell(row=i, column=3, value=meaning)

    start = 6 + len(summary_rows) + 1
    ws.cell(row=start, column=1, value='Validation Status').font = TITLE_FONT
    _write_headers(ws, ['Status', 'Citations', 'Meaning'], row=start + 1)
    status_meaning = {
        'Valid': 'Checked and loading normally.',
        'Broken': 'Checked and returning an HTTP error.',
        'Blocked': 'The page refused our crawler; the page itself may be fine.',
        'Pending': 'Not yet checked — validation has not run for these URLs.',
    }
    for i, label in enumerate(['Valid', 'Broken', 'Blocked', 'Pending'], start=start + 2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=status_counts.get(label, 0))
        ws.cell(row=i, column=3, value=status_meaning[label])
    _autosize(ws, max_width=80)

    # ---------- Sheet 2: Citations ----------
    ws = wb.create_sheet('Citations')
    headers = [
        '#', 'Status', 'HTTP Code', 'Source URL', 'Source Domain', 'Type',
        'Platform', 'Position In Answer', 'Prompt', 'Prompt Group',
        'Cited On', 'Last Checked',
    ]
    _write_headers(ws, headers)
    for row in rows:
        ws.append(row)
    for col in ('K', 'L'):
        for cell in ws[col][1:]:
            cell.number_format = 'yyyy-mm-dd hh:mm'
    _finish_table(ws, len(rows), len(headers))

    # ---------- Sheet 3: By Source ----------
    ws = wb.create_sheet('By Source')
    headers = ['Source Domain', 'Citations', 'Share %', 'Type', 'Platforms', 'Last Cited']
    _write_headers(ws, headers)
    ranked = sorted(by_source.items(), key=lambda kv: kv[1]['count'], reverse=True)
    for host, data in ranked:
        ws.append([
            host,
            data['count'],
            round(data['count'] / total * 100, 1) if total else 0,
            'Your Domain' if data['own'] else 'Third Party',
            ', '.join(sorted(data['platforms'])),
            _naive(data['last']),
        ])
    for cell in ws['F'][1:]:
        cell.number_format = 'yyyy-mm-dd hh:mm'
    _finish_table(ws, len(ranked), len(headers))

    # ---------- Sheet 4: By Platform ----------
    ws = wb.create_sheet('By Platform')
    headers = ['Platform', 'Citations', 'Share %']
    _write_headers(ws, headers)
    platform_rows = sorted(by_platform.items(), key=lambda kv: kv[1], reverse=True)
    for platform, count in platform_rows:
        ws.append([platform, count, round(count / total * 100, 1) if total else 0])
    _finish_table(ws, len(platform_rows), len(headers))

    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    safe_name = ''.join(ch for ch in (domain.name or 'domain') if ch.isalnum() or ch in ' -_').strip().replace(' ', '_')
    filename = f'citations_{safe_name}_{stamp}.xlsx'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    logger.info(f"Citations export: domain {domain_id}, {total} citations, {len(by_source)} sources")
    return response
