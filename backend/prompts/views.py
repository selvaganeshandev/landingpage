from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response
from rest_framework import status
from rest_framework.permissions import IsAuthenticated, AllowAny
from django.db.models import Q, Count, Avg, F, Sum
from django.shortcuts import get_object_or_404
from django.utils import timezone
from datetime import timedelta, datetime
from .models import PromptGroup, Prompt, PromptAnalytics, PromptGroupMetricSnapshot, PromptMetricSnapshot, DomainMetricSnapshot
from domains.models import Domain, DomainAccess
from .serializers import PromptAnalyticsSerializer, PromptGroupSerializer, PromptSerializer
import json
import logging
import statistics
from dateutil.relativedelta import relativedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
import io

logger = logging.getLogger(__name__)

# Import OpenAI client helper
def get_openai_client():
    """Return OpenAI client if configured in Django settings; else raise."""
    from django.conf import settings
    api_key = getattr(settings, "OPENAI_API_KEY", None)
    if not api_key:
        raise Exception("OpenAI API key not configured")
    try:
        from openai import OpenAI  # lazy import
        return OpenAI(api_key=api_key, timeout=60)
    except Exception as e:
        raise Exception(f"Failed to initialize OpenAI client: {e}")


@api_view(['GET'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def get_mentions(request):
    """
    Get mentions from PromptAnalytics where is_mention=True and is_published=True
    """
    # Require domain_id and access check
    domain_id = request.GET.get('domain_id')
    if not domain_id:
        return Response({
            'mentions': [],
            'total_count': 0,
            'filters_applied': {
                'search': request.GET.get('search', ''),
                'platform': request.GET.get('platform', ''),
                'sentiment': request.GET.get('sentiment', ''),
                'domain_id': None,
            },
            'available_platforms': ['ChatGPT', 'Google Gemini', 'Perplexity'],
            'available_sentiments': ['Positive', 'Negative', 'Neutral']
        })

    user = getattr(request, 'user', None)
    has_access = False
    if user and getattr(user, 'is_authenticated', False):
        if getattr(user, 'role', '') == 'super_admin':
            has_access = True
        else:
            has_access = DomainAccess.objects.filter(user=user, domain_id=domain_id).exists()
    if not has_access:
        return Response({'error': 'Forbidden: no access to this domain.'}, status=status.HTTP_403_FORBIDDEN)

    # Get analytics records - show all or only mentions based on show_all parameter
    show_all = request.GET.get('show_all', 'false').lower() == 'true'
    if show_all:
        # Show all prompt analytics (published) for this domain
        mentions = PromptAnalytics.objects.filter(is_published=True, prompt__group__domain_id=domain_id)
    else:
        # Show only mentions
        mentions = PromptAnalytics.objects.filter(is_mention=True, is_published=True, prompt__group__domain_id=domain_id)
    
    # Apply organization filter if user is authenticated and not superuser
    # Temporarily skip organization filtering for testing
    # if hasattr(request, 'user') and request.user.is_authenticated and not request.user.is_superuser:
    #     mentions = mentions.filter()
    
    # Apply search filter if provided
    search_query = request.GET.get('search', '')
    if search_query:
        mentions = mentions.filter(
            Q(prompt__prompt__icontains=search_query) |
            Q(context_summary__icontains=search_query) |
            Q(prompt__group__domain__name__icontains=search_query)
        )
    
    # Apply platform filter if provided
    platform = request.GET.get('platform', '')
    if platform and platform != 'all' and platform != 'All Platforms':
        mentions = mentions.filter(platform__icontains=platform)
    
    # Apply sentiment filter if provided
    sentiment = request.GET.get('sentiment', '')
    if sentiment and sentiment.lower() != 'all' and sentiment.lower() != 'all sentiments':
        # Model field is sentiment_category in shared schema
        mentions = mentions.filter(sentiment_category__iexact=sentiment)
    
    # Order by most recent first
    mentions = mentions.order_by('-created_at')

    # Pagination
    try:
        limit = int(request.GET.get('limit', 20))
        offset = int(request.GET.get('offset', 0))
    except ValueError:
        limit = 20
        offset = 0
    limit = max(1, min(limit, 100))
    offset = max(0, offset)
    total_count = mentions.count()
    mentions = mentions[offset:offset + limit]
    
    # Helper function to extract domain name from URL
    def _extract_domain_name(url):
        """Extract a readable domain name from URL"""
        try:
            from urllib.parse import urlparse
            parsed = urlparse(url)
            domain = parsed.netloc or parsed.path
            if domain.startswith('www.'):
                domain = domain[4:]
            # Capitalize and format
            parts = domain.split('.')
            if len(parts) >= 2:
                return parts[0].capitalize() + ' ' + parts[1].capitalize()
            return domain.capitalize()
        except:
            return 'Source'
    
    # Helper function to extract headline from context around URL
    def _extract_headline_from_context(url, context_summary):
        """Extract a headline/quote from context around the URL"""
        if not context_summary or not url:
            return None
        try:
            import re
            # Find the URL in the context
            url_lower = url.lower()
            context_lower = context_summary.lower()
            idx = context_lower.find(url_lower)
            if idx == -1:
                return None
            
            # Extract text around the URL (300 chars before and after)
            start = max(0, idx - 300)
            end = min(len(context_summary), idx + len(url) + 300)
            snippet = context_summary[start:end]
            url_pos_in_snippet = snippet.lower().find(url_lower)
            
            if url_pos_in_snippet == -1:
                return None
            
            # Get text before the URL
            text_before = snippet[:url_pos_in_snippet].strip()
            
            # Strategy 1: Look for quoted text (text in quotes)
            quoted_matches = re.findall(r'["\']([^"\']{10,100})["\']', text_before)
            if quoted_matches:
                # Return the last (most recent) quote
                return quoted_matches[-1].strip()
            
            # Strategy 2: Look for sentences ending with punctuation before the URL
            # Split by sentence endings
            sentences = re.split(r'[.!?]\s+', text_before)
            if sentences:
                # Get the last complete sentence before the URL
                last_sentence = sentences[-1].strip()
                # Clean up common prefixes
                last_sentence = re.sub(r'^(For|As|According to|In|On|The|A|An)\s+', '', last_sentence, flags=re.IGNORECASE)
                if len(last_sentence) > 10 and len(last_sentence) < 150:
                    return last_sentence.strip()
            
            # Strategy 3: Extract key phrases (look for patterns like "stands out", "top choice", etc.)
            # Look for common patterns that indicate key statements
            patterns = [
                r'([^.!?]{0,50}(?:stands out|top choice|best|superior|excellent|outstanding|recommended|highly rated)[^.!?]{0,50})',
                r'([^.!?]{0,50}(?:key benefits|features|advantages|benefits)[^.!?]{0,50})',
                r'([^.!?]{0,50}(?:notable|significant|important|noteworthy)[^.!?]{0,50})',
            ]
            for pattern in patterns:
                matches = re.findall(pattern, text_before, re.IGNORECASE)
                if matches:
                    phrase = matches[-1].strip()
                    if len(phrase) > 10 and len(phrase) < 150:
                        return phrase.strip()
            
            # Strategy 4: Extract last 5-15 words as fallback
            words = text_before.split()
            if len(words) >= 5:
                # Get last 5-15 words
                phrase = ' '.join(words[-15:]).strip()
                # Remove common prefixes
                phrase = re.sub(r'^(For|As|According to|In|On|The|A|An)\s+', '', phrase, flags=re.IGNORECASE)
                if len(phrase) > 10 and len(phrase) < 150:
                    return phrase.strip()
        except Exception as e:
            logger.warning(f"Error extracting headline from context: {e}")
            pass
        return None
    
    # Prepare response data
    mentions_data = []
    for mention in mentions:
        # Extract and format citations from citation_list
        citations_data = []
        citation_list = getattr(mention, 'citation_list', None) or []
        context_summary = mention.context_summary or ''
        
        if citation_list and isinstance(citation_list, list):
            for idx, citation_url in enumerate(citation_list):
                if not citation_url or not isinstance(citation_url, str):
                    continue
                
                # Extract headline from context around this URL
                headline = _extract_headline_from_context(citation_url, context_summary)
                
                # Extract source name from URL
                source_name = _extract_domain_name(citation_url)
                
                # Generate description based on source type
                url_lower = citation_url.lower()
                if 'product' in url_lower or 'shop' in url_lower or 'store' in url_lower:
                    description = 'Product page citing key benefits and features'
                elif 'review' in url_lower or 'rating' in url_lower:
                    description = 'Third-party review and analysis'
                elif 'blog' in url_lower or 'article' in url_lower:
                    description = 'Article discussing key features'
                elif 'healthline' in url_lower or 'medical' in url_lower or 'health' in url_lower:
                    description = 'Third-party nutritional analysis'
                elif 'official' in url_lower or 'site' in url_lower:
                    description = 'Official site with product information'
                else:
                    description = 'Source providing relevant information'
                
                citations_data.append({
                    'id': idx + 1,
                    'text': headline or f'Citation from {source_name}',
                    'source': source_name,
                    'url': citation_url,
                    'description': description
                })
        
        mention_data = {
            'id': mention.id,
            'rank': len(mentions_data) + 1,  # Simple ranking based on order
            'mention_text_short': mention.prompt.prompt[:50] + '...' if len(mention.prompt.prompt) > 50 else mention.prompt.prompt,
            'mention_text_long': mention.prompt.prompt,
            'description': (mention.context_summary[:500] + '...') if mention.context_summary and len(mention.context_summary) > 500 else (mention.context_summary or ''),
            'platform': mention.platform,
            'sentiment': getattr(mention, 'sentiment_category', None) or getattr(mention, 'sentiment', ''),
            'sentiment_score': float(mention.sentiment_score),
            'total_mentions': mention.total_mentions,
            'total_citations': mention.total_citations,
            'position': float(mention.position),
            'timestamp': mention.created_at.isoformat(),
            'time_ago': _get_time_ago(mention.created_at),
            'domain_name': (mention.prompt.group.domain.name if mention.prompt and mention.prompt.group else ''),
            'domain_url': (mention.prompt.group.domain.url if mention.prompt and mention.prompt.group else ''),
            'group_id': mention.prompt.group.group_id if mention.prompt.group else None,
            'track_status': mention.prompt.track_status,
            'type': mention.prompt.type,
            'citations': citations_data,
            'citations_count': len(citations_data),
            'views': mention.views,
            'shares': mention.shares,
            'engagement_score': float(mention.engagement_score),
            'competitor_mentions': (getattr(mention, 'competitor_mention_list', None) or getattr(mention, 'competitor_mentions', []) or []),
            'key_topics': (getattr(mention, 'topic_list', None) or getattr(mention, 'key_topics', []) or [])
        }
        mentions_data.append(mention_data)
    
    return Response({
        'mentions': mentions_data,
        'total_count': total_count,
        'filters_applied': {
            'search': search_query,
            'platform': platform,
            'sentiment': sentiment,
            'domain_id': domain_id
        },
        'pagination': {
            'limit': limit,
            'offset': offset,
            'returned': len(mentions_data)
        },
        'available_platforms': _get_available_platforms(),
        'available_sentiments': _get_available_sentiments()
    })


def _get_time_ago(created_at):
    """Helper function to get human-readable time ago"""
    from django.utils import timezone
    from datetime import timedelta
    
    now = timezone.now()
    diff = now - created_at
    
    if diff.days > 0:
        return f"{diff.days} day{'s' if diff.days > 1 else ''} ago"
    elif diff.seconds > 3600:
        hours = diff.seconds // 3600
        return f"{hours} hour{'s' if hours > 1 else ''} ago"
    elif diff.seconds > 60:
        minutes = diff.seconds // 60
        return f"{minutes} minute{'s' if minutes > 1 else ''} ago"
    else:
        return "Just now"


def _get_available_platforms():
    """Get list of available platforms for filtering"""
    platforms = PromptAnalytics.objects.filter(is_mention=True, is_published=True).values_list('platform', flat=True).distinct()
    return list(platforms)


def _get_available_sentiments():
    """Get list of available sentiments for filtering"""
    sentiments = PromptAnalytics.objects.filter(is_mention=True, is_published=True).values_list('sentiment_category', flat=True).distinct()
    return list(sentiments)


@api_view(['GET'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def get_mention_filters(request):
    """
    Get available filter options for mentions
    """
    # Apply organization filter if user is not superuser
    base_query = PromptAnalytics.objects.filter(is_mention=True, is_published=True)
    # Temporarily skip organization filtering for testing
    # if not request.user.is_superuser:
    #     base_query = base_query.filter()
    
    # Get available platforms
    platforms = list(base_query.values_list('platform', flat=True).distinct())
    
    # Get available sentiments
    sentiments = list(base_query.values_list('sentiment_category', flat=True).distinct())
    
    # Get total counts by platform
    platform_counts = {}
    for platform in platforms:
        count = base_query.filter(platform=platform).count()
        platform_counts[platform] = count
    
    # Get total counts by sentiment
    sentiment_counts = {}
    for sentiment in sentiments:
        count = base_query.filter(sentiment_category=sentiment).count()
        sentiment_counts[sentiment] = count
    
    return Response({
        'platforms': [
            {'name': platform, 'count': platform_counts.get(platform, 0)}
            for platform in platforms
        ],
        'sentiments': [
            {'name': sentiment, 'count': sentiment_counts.get(sentiment, 0)}
            for sentiment in sentiments
        ],
        'total_mentions': base_query.count()
    })


@api_view(['GET'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def get_mention_detail(request, analytics_id):
    """
    Get detailed information for a specific mention (PromptAnalytics record)
    """
    try:
        # Get the specific analytics record
        analytics_record = get_object_or_404(PromptAnalytics, id=analytics_id)
        
        # Apply organization filter if user is not superuser
        # Temporarily disabled for testing
        # if not request.user.is_superuser:
        #     if analytics_record.organisation != request.user.organisation:
        #         return Response(
        #             {'error': 'You do not have permission to access this mention.'},
        #             status=status.HTTP_403_FORBIDDEN
        #         )
        
        # Get related objects
        prompt = analytics_record.prompt
        group = prompt.group if prompt.group else None
        
        # Use the same citation extraction logic as get_mentions
        # Helper function to extract domain name from URL
        def _extract_domain_name(url):
            """Extract a readable domain name from URL"""
            try:
                from urllib.parse import urlparse
                parsed = urlparse(url)
                domain = parsed.netloc or parsed.path
                if domain.startswith('www.'):
                    domain = domain[4:]
                # Capitalize and format
                parts = domain.split('.')
                if len(parts) >= 2:
                    return parts[0].capitalize() + ' ' + parts[1].capitalize()
                return domain.capitalize()
            except:
                return 'Source'
        
        # Helper function to extract headline from context around URL
        def _extract_headline_from_context(url, context_summary):
            """Extract a headline/quote from context around the URL"""
            if not context_summary or not url:
                return None
            try:
                import re
                # Find the URL in the context
                url_lower = url.lower()
                context_lower = context_summary.lower()
                idx = context_lower.find(url_lower)
                if idx == -1:
                    return None
                
                # Extract text around the URL (300 chars before and after)
                start = max(0, idx - 300)
                end = min(len(context_summary), idx + len(url) + 300)
                snippet = context_summary[start:end]
                url_pos_in_snippet = snippet.lower().find(url_lower)
                
                if url_pos_in_snippet == -1:
                    return None
                
                # Get text before the URL
                text_before = snippet[:url_pos_in_snippet].strip()
                
                # Strategy 1: Look for quoted text (text in quotes)
                quoted_matches = re.findall(r'["\']([^"\']{10,100})["\']', text_before)
                if quoted_matches:
                    # Return the last (most recent) quote
                    return quoted_matches[-1].strip()
                
                # Strategy 2: Look for sentences ending with punctuation before the URL
                # Split by sentence endings
                sentences = re.split(r'[.!?]\s+', text_before)
                if sentences:
                    # Get the last complete sentence before the URL
                    last_sentence = sentences[-1].strip()
                    # Clean up common prefixes
                    last_sentence = re.sub(r'^(For|As|According to|In|On|The|A|An)\s+', '', last_sentence, flags=re.IGNORECASE)
                    if len(last_sentence) > 10 and len(last_sentence) < 150:
                        return last_sentence.strip()
                
                # Strategy 3: Extract key phrases (look for patterns like "stands out", "top choice", etc.)
                # Look for common patterns that indicate key statements
                patterns = [
                    r'([^.!?]{0,50}(?:stands out|top choice|best|superior|excellent|outstanding|recommended|highly rated)[^.!?]{0,50})',
                    r'([^.!?]{0,50}(?:key benefits|features|advantages|benefits)[^.!?]{0,50})',
                    r'([^.!?]{0,50}(?:notable|significant|important|noteworthy)[^.!?]{0,50})',
                ]
                for pattern in patterns:
                    matches = re.findall(pattern, text_before, re.IGNORECASE)
                    if matches:
                        phrase = matches[-1].strip()
                        if len(phrase) > 10 and len(phrase) < 150:
                            return phrase.strip()
                
                # Strategy 4: Extract last 5-15 words as fallback
                words = text_before.split()
                if len(words) >= 5:
                    # Get last 5-15 words
                    phrase = ' '.join(words[-15:]).strip()
                    # Remove common prefixes
                    phrase = re.sub(r'^(For|As|According to|In|On|The|A|An)\s+', '', phrase, flags=re.IGNORECASE)
                    if len(phrase) > 10 and len(phrase) < 150:
                        return phrase.strip()
            except Exception as e:
                logger.warning(f"Error extracting headline from context: {e}")
                pass
            return None
        
        # Extract and format citations from citation_list
        citations_data = []
        citation_list = getattr(analytics_record, 'citation_list', None) or []
        context_summary = analytics_record.context_summary or ''
        
        if citation_list and isinstance(citation_list, list):
            for idx, citation_url in enumerate(citation_list):
                if not citation_url or not isinstance(citation_url, str):
                    continue
                
                # Extract headline from context around this URL
                headline = _extract_headline_from_context(citation_url, context_summary)
                
                # Extract source name from URL
                source_name = _extract_domain_name(citation_url)
                
                # Generate description based on source type
                url_lower = citation_url.lower()
                if 'product' in url_lower or 'shop' in url_lower or 'store' in url_lower:
                    description = 'Product page citing key benefits and features'
                elif 'review' in url_lower or 'rating' in url_lower:
                    description = 'Third-party review and analysis'
                elif 'blog' in url_lower or 'article' in url_lower:
                    description = 'Article discussing key features'
                elif 'healthline' in url_lower or 'medical' in url_lower or 'health' in url_lower:
                    description = 'Third-party nutritional analysis'
                elif 'official' in url_lower or 'site' in url_lower:
                    description = 'Official site with product information'
                else:
                    description = 'Source providing relevant information'
                
                citations_data.append({
                    'id': idx + 1,
                    'text': headline or f'Citation from {source_name}',
                    'source': source_name,
                    'source_name': source_name,  # Keep for backward compatibility
                    'url': citation_url,
                    'source_url': citation_url,  # Keep for backward compatibility
                    'description': description,
                    'reliability': 'Verified',  # Default reliability
                    'referenced_at': analytics_record.created_at.isoformat()
                })
        
        # Prepare comprehensive response data
        response_data = {
            # Basic mention info
            'id': analytics_record.id,
            'rank': 1,  # This would need to be calculated based on position/score
            'platform': analytics_record.platform,
            'sentiment': getattr(analytics_record, 'sentiment_category', 'neutral'),
            'sentiment_score': float(analytics_record.sentiment_score),
            'created_at': analytics_record.created_at.isoformat(),
            'time_ago': _get_time_ago(analytics_record.created_at),
            
            # Prompt details
            'prompt_text': prompt.prompt,
            'prompt_text_short': prompt.prompt[:100] + '...' if len(prompt.prompt) > 100 else prompt.prompt,
            'full_ai_response': analytics_record.context_summary or '',
            
            # Metrics
            'total_mentions': analytics_record.total_mentions,
            'total_citations': analytics_record.total_citations,
            'position': float(analytics_record.position),
            'is_mention': analytics_record.is_mention,
            
            # Tracking info
            'track_status': prompt.track_status,
            'type': prompt.type,
            'last_tracked_at': prompt.tracked_at.isoformat() if prompt.tracked_at else None,
            'track_message': prompt.track_message,
            
            # Domain and group info
            'domain_name': (group.domain.name if group else ''),
            'domain_url': (group.domain.url if group else ''),
            'group_id': group.group_id if group else None,
            'prompt_id': prompt.id,
            
            # Citations with detailed structure
            'citations': citations_data,
            'citations_count': len(citations_data),
            
            # Engagement metrics
            'views': getattr(analytics_record, 'views', None),
            'shares': getattr(analytics_record, 'shares', None),
            
            # Historical data (placeholder)
            'position_trend': [],
            'key_topics': (getattr(analytics_record, 'topic_list', None) or getattr(analytics_record, 'key_topics', []) or []),
            'competitor_mentions': (getattr(analytics_record, 'competitor_mention_list', None) or getattr(analytics_record, 'competitor_mentions', []) or []),
        }
        
        return Response(response_data)
        
    except Exception as e:
        return Response(
            {'error': f'Failed to retrieve mention details: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_mention_trends(request):
    """
    Get mention trends over time for analytics
    """
    try:
        # Get date range from query params
        days = int(request.GET.get('days', 30))
        end_date = timezone.now()
        start_date = end_date - timedelta(days=days)
        
        # Get mentions for the organization
        mentions = PromptAnalytics.objects.filter(
            is_mention=True,
            is_published=True,
            # organisation__isnull=True,  # Temporarily disabled for testing
            created_at__range=[start_date, end_date]
        )
        
        # Group by date and platform
        trends_data = []
        current_date = start_date
        
        while current_date <= end_date:
            date_str = current_date.strftime('%Y-%m-%d')
            day_mentions = mentions.filter(created_at__date=current_date.date())
            
            platform_data = {}
            for platform in ['ChatGPT', 'Claude', 'Perplexity', 'Gemini', 'Grok']:
                platform_mentions = day_mentions.filter(platform=platform)
                platform_data[platform.lower()] = {
                    'count': platform_mentions.count(),
                    'avg_position': float(platform_mentions.aggregate(avg_pos=Avg('position'))['avg_pos'] or 0),
                    'avg_sentiment': float(platform_mentions.aggregate(avg_sent=Avg('sentiment_score'))['avg_sent'] or 0)
                }
            
            trends_data.append({
                'date': date_str,
                'total_mentions': day_mentions.count(),
                'avg_position': float(day_mentions.aggregate(avg_pos=Avg('position'))['avg_pos'] or 0),
                'platforms': platform_data
            })
            
            current_date += timedelta(days=1)
        
        return Response({
            'trends': trends_data,
            'period': f'{days} days',
            'total_mentions': mentions.count()
        })
        
    except Exception as e:
        return Response(
            {'error': f'Failed to retrieve mention trends: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_mention_analytics(request):
    """
    Get comprehensive analytics for mentions
    """
    try:
        # Get mentions for the organization
        mentions = PromptAnalytics.objects.filter(
            is_mention=True,
            is_published=True
        )
        
        # Platform breakdown
        platform_stats = mentions.values('platform').annotate(
            count=Count('id'),
            avg_position=Avg('position'),
            avg_sentiment=Avg('sentiment_score'),
            total_views=Count('views'),
            total_shares=Count('shares')
        ).order_by('-count')
        
        # Sentiment breakdown
        sentiment_stats = mentions.values('sentiment_category').annotate(
            count=Count('id'),
            avg_position=Avg('position'),
            avg_sentiment_score=Avg('sentiment_score')
        ).order_by('-count')
        
        # Top performing mentions
        top_mentions = mentions.order_by('-engagement_score')[:10]
        top_mentions_data = []
        for mention in top_mentions:
            top_mentions_data.append({
                'id': mention.id,
                'prompt_text': mention.prompt.prompt[:100] + '...' if len(mention.prompt.prompt) > 100 else mention.prompt.prompt,
                'platform': mention.platform,
                'position': float(mention.position),
                'sentiment': getattr(mention, 'sentiment_category', None) or getattr(mention, 'sentiment', ''),
                'engagement_score': float(mention.engagement_score),
                'views': mention.views,
                'shares': mention.shares,
                'created_at': mention.created_at.isoformat()
            })
        
        # Competitor analysis
        competitor_data = {}
        for mention in mentions:
            comp_list = getattr(mention, 'competitor_mention_list', None) or getattr(mention, 'competitor_mentions', []) or []
            if comp_list:
                for competitor in comp_list:
                    if competitor not in competitor_data:
                        competitor_data[competitor] = {'count': 0, 'mentions': []}
                    competitor_data[competitor]['count'] += 1
                    competitor_data[competitor]['mentions'].append({
                        'id': mention.id,
                        'platform': mention.platform,
                        'sentiment': getattr(mention, 'sentiment_category', None) or getattr(mention, 'sentiment', ''),
                        'position': float(mention.position)
                    })
        
        return Response({
            'platform_stats': list(platform_stats),
            'sentiment_stats': list(sentiment_stats),
            'top_mentions': top_mentions_data,
            'competitor_analysis': competitor_data,
            'total_mentions': mentions.count(),
            'avg_position': float(mentions.aggregate(avg_pos=Avg('position'))['avg_pos'] or 0),
            'avg_sentiment': float(mentions.aggregate(avg_sent=Avg('sentiment_score'))['avg_sent'] or 0)
        })
        
    except Exception as e:
        return Response(
            {'error': f'Failed to retrieve mention analytics: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['GET'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def get_related_mentions(request, analytics_id):
    """
    Get related mentions for a specific mention
    """
    try:
        # Get the original mention
        original_mention = get_object_or_404(PromptAnalytics, id=analytics_id)
        
        # Access checks can be added here if needed (organisation removed)
        
        # Find related mentions based on similar prompts or same group
        related_mentions = PromptAnalytics.objects.filter(
            is_mention=True,
            is_published=True
        ).exclude(id=analytics_id)
        
        # Filter by same group first (only if the original mention has a group)
        same_group_mentions = []
        if original_mention.prompt.group:
            same_group_mentions = related_mentions.filter(
                prompt__group=original_mention.prompt.group
            ).order_by('-created_at')[:5]
        
        # Filter by similar platform
        same_platform_mentions = related_mentions.filter(
            platform=original_mention.platform
        )
        if original_mention.prompt.group:
            same_platform_mentions = same_platform_mentions.exclude(prompt__group=original_mention.prompt.group)
        same_platform_mentions = same_platform_mentions.order_by('-created_at')[:3]
        
        # Combine and format results
        related_data = []
        for mention in same_group_mentions:
            related_data.append({
                'id': mention.id,
                'prompt_text': mention.prompt.prompt[:100] + '...' if len(mention.prompt.prompt) > 100 else mention.prompt.prompt,
                'platform': mention.platform,
                'position': float(mention.position),
                'sentiment': getattr(mention, 'sentiment_category', None) or getattr(mention, 'sentiment', ''),
                'created_at': mention.created_at.isoformat(),
                'time_ago': _get_time_ago(mention.created_at),
                'relation_type': 'same_group'
            })
        
        for mention in same_platform_mentions:
            related_data.append({
                'id': mention.id,
                'prompt_text': mention.prompt.prompt[:100] + '...' if len(mention.prompt.prompt) > 100 else mention.prompt.prompt,
                'platform': mention.platform,
                'position': float(mention.position),
                'sentiment': getattr(mention, 'sentiment_category', None) or getattr(mention, 'sentiment', ''),
                'created_at': mention.created_at.isoformat(),
                'time_ago': _get_time_ago(mention.created_at),
                'relation_type': 'same_platform'
            })
        
        return Response({
            'related_mentions': related_data,
            'original_mention': {
                'id': original_mention.id,
                'prompt_text': original_mention.prompt.prompt,
                'platform': original_mention.platform,
                'group_id': original_mention.prompt.group.group_id if original_mention.prompt.group else None
            }
        })
        
    except Exception as e:
        return Response(
            {'error': f'Failed to retrieve related mentions: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def export_mentions(request):
    """
    Export mentions data
    """
    try:
        # Get filter parameters
        platform = request.data.get('platform', '')
        sentiment = request.data.get('sentiment', '')
        date_from = request.data.get('date_from', '')
        date_to = request.data.get('date_to', '')
        export_format = request.data.get('format', 'csv')
        
        # Build query
        mentions = PromptAnalytics.objects.filter(
            is_mention=True,
            is_published=True
        )
        
        if platform and platform != 'all':
            mentions = mentions.filter(platform__icontains=platform)
        
        if sentiment and sentiment != 'all':
            mentions = mentions.filter(sentiment_category=sentiment)
        
        if date_from:
            mentions = mentions.filter(created_at__gte=date_from)
        
        if date_to:
            mentions = mentions.filter(created_at__lte=date_to)
        
        # Prepare export data
        export_data = []
        for mention in mentions:
            export_data.append({
                'id': mention.id,
                'prompt': mention.prompt.prompt,
                'platform': mention.platform,
                'sentiment': getattr(mention, 'sentiment_category', None) or getattr(mention, 'sentiment', ''),
                'position': float(mention.position),
                'total_mentions': mention.total_mentions,
                'total_citations': mention.total_citations,
                'views': mention.views,
                'shares': mention.shares,
                'engagement_score': float(mention.engagement_score),
                'created_at': mention.created_at.isoformat(),
                'domain_name': (mention.prompt.group.domain.name if mention.prompt and mention.prompt.group else ''),
                'group_id': mention.prompt.group.group_id if mention.prompt and mention.prompt.group else None
            })
        
        # For now, return the data (in real implementation, you'd generate actual file)
        return Response({
            'message': 'Export data prepared',
            'format': export_format,
            'count': len(export_data),
            'data': export_data,
            'download_url': f'/api/prompts/mentions/export/download/{len(export_data)}/'  # Placeholder
        })
        
    except Exception as e:
        return Response(
            {'error': f'Failed to export mentions: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


def _format_mention_for_export(analytics_record):
    """
    Format a PromptAnalytics record into the export format
    Returns a dictionary matching the specified JSON structure
    """
    prompt = analytics_record.prompt
    group = prompt.group if prompt.group else None
    
    # Helper function to extract domain name from URL
    def _extract_domain_name(url):
        """Extract a readable domain name from URL"""
        try:
            from urllib.parse import urlparse
            parsed = urlparse(url)
            domain = parsed.netloc or parsed.path
            if domain.startswith('www.'):
                domain = domain[4:]
            # Capitalize and format
            parts = domain.split('.')
            if len(parts) >= 2:
                return parts[0].capitalize() + ' ' + parts[1].capitalize()
            return domain.capitalize()
        except:
            return 'Source'
    
    # Helper function to extract headline from context around URL
    def _extract_headline_from_context(url, context_summary):
        """Extract a headline/quote from context around the URL"""
        if not context_summary or not url:
            return None
        try:
            import re
            # Find the URL in the context
            url_lower = url.lower()
            context_lower = context_summary.lower()
            idx = context_lower.find(url_lower)
            if idx == -1:
                return None
            
            # Extract text around the URL (300 chars before and after)
            start = max(0, idx - 300)
            end = min(len(context_summary), idx + len(url) + 300)
            snippet = context_summary[start:end]
            url_pos_in_snippet = snippet.lower().find(url_lower)
            
            if url_pos_in_snippet == -1:
                return None
            
            # Get text before the URL
            text_before = snippet[:url_pos_in_snippet].strip()
            
            # Strategy 1: Look for quoted text (text in quotes)
            quoted_matches = re.findall(r'["\']([^"\']{10,100})["\']', text_before)
            if quoted_matches:
                return quoted_matches[-1].strip()
            
            # Strategy 2: Look for sentences ending with punctuation before the URL
            sentences = re.split(r'[.!?]\s+', text_before)
            if sentences:
                last_sentence = sentences[-1].strip()
                last_sentence = re.sub(r'^(For|As|According to|In|On|The|A|An)\s+', '', last_sentence, flags=re.IGNORECASE)
                if len(last_sentence) > 10 and len(last_sentence) < 150:
                    return last_sentence.strip()
            
            # Strategy 3: Extract last 5-15 words as fallback
            words = text_before.split()
            if len(words) >= 5:
                phrase = ' '.join(words[-15:]).strip()
                phrase = re.sub(r'^(For|As|According to|In|On|The|A|An)\s+', '', phrase, flags=re.IGNORECASE)
                if len(phrase) > 10 and len(phrase) < 150:
                    return phrase.strip()
        except Exception as e:
            logger.warning(f"Error extracting headline from context: {e}")
        return None
    
    # Extract and format citations from citation_list
    citations_data = []
    citation_list = getattr(analytics_record, 'citation_list', None) or []
    context_summary = analytics_record.context_summary or ''
    
    if citation_list and isinstance(citation_list, list):
        for idx, citation_url in enumerate(citation_list):
            if not citation_url or not isinstance(citation_url, str):
                continue
            
            # Extract headline from context around this URL
            headline = _extract_headline_from_context(citation_url, context_summary)
            
            # Extract source name from URL
            source_name = _extract_domain_name(citation_url)
            
            # Generate description based on source type
            url_lower = citation_url.lower()
            if 'product' in url_lower or 'shop' in url_lower or 'store' in url_lower:
                description = 'Product page citing key benefits and features'
            elif 'review' in url_lower or 'rating' in url_lower:
                description = 'Third-party review and analysis'
            elif 'blog' in url_lower or 'article' in url_lower:
                description = 'Article discussing key features'
            elif 'healthline' in url_lower or 'medical' in url_lower or 'health' in url_lower:
                description = 'Third-party nutritional analysis'
            elif 'official' in url_lower or 'site' in url_lower:
                description = 'Official site with product information'
            else:
                description = 'Source providing relevant information'
            
            citations_data.append({
                'id': idx + 1,
                'text': headline or f'Citation from {source_name}',
                'source': source_name,
                'source_name': source_name,
                'url': citation_url,
                'source_url': citation_url,
                'description': description,
                'reliability': 'Verified',
                'referenced_at': analytics_record.created_at.isoformat()
            })
    
    # Format the mention data according to the specified structure
    formatted_data = {
        'mention_id': analytics_record.id,
        'platform': analytics_record.platform,
        'sentiment': getattr(analytics_record, 'sentiment_category', 'neutral'),
        'sentiment_score': float(analytics_record.sentiment_score),
        'prompt_text': prompt.prompt,
        'full_ai_response': analytics_record.context_summary or '',
        'total_mentions': analytics_record.total_mentions,
        'total_citations': analytics_record.total_citations,
        'position': float(analytics_record.position),
        'created_at': analytics_record.created_at.isoformat(),
        'domain_name': (group.domain.name if group else ''),
        'citations': citations_data,
        'key_topics': (getattr(analytics_record, 'topic_list', None) or getattr(analytics_record, 'key_topics', []) or [])
    }
    
    return formatted_data


def _create_excel_from_mentions(mentions_data, sheet_name='Mentions'):
    """
    Create an Excel workbook from mentions data
    Returns a BytesIO buffer with the Excel file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    # Define styles
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Headers
    headers = [
        'Mention ID', 'Platform', 'Sentiment', 'Sentiment Score', 'Prompt Text',
        'Full AI Response', 'Total Mentions', 'Total Citations', 'Position',
        'Created At', 'Domain Name', 'Citations (JSON)', 'Key Topics (JSON)'
    ]
    
    # Write headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Write data
    logger.info(f"Writing {len(mentions_data)} mentions to Excel")
    for row_num, mention in enumerate(mentions_data, 2):
        # Convert citations and key_topics to JSON strings for Excel
        citations_json = json.dumps(mention.get('citations', []), indent=2) if mention.get('citations') else ''
        key_topics_json = json.dumps(mention.get('key_topics', []), indent=2) if mention.get('key_topics') else ''
        
        data_row = [
            mention.get('mention_id', ''),
            mention.get('platform', ''),
            mention.get('sentiment', ''),
            mention.get('sentiment_score', 0),
            mention.get('prompt_text', ''),
            mention.get('full_ai_response', ''),
            mention.get('total_mentions', 0),
            mention.get('total_citations', 0),
            mention.get('position', 0),
            mention.get('created_at', ''),
            mention.get('domain_name', ''),
            citations_json,
            key_topics_json
        ]
        
        for col_num, value in enumerate(data_row, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = value
            cell.border = border
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    
    logger.info(f"Excel file created with {len(mentions_data)} rows of data")
    
    # Auto-adjust column widths
    for col_num in range(1, len(headers) + 1):
        column_letter = get_column_letter(col_num)
        max_length = 0
        for row in ws[column_letter]:
            try:
                if len(str(row.value)) > max_length:
                    max_length = len(str(row.value))
            except:
                pass
        # Set width with some padding, but cap at 50 for very long content
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@api_view(['GET'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def export_mentions_list(request):
    """
    Export all mentions to Excel file
    Supports the same filters as get_mentions
    """
    try:
        # Get filter parameters (same as get_mentions)
        domain_id = request.GET.get('domain_id')
        if not domain_id:
            return Response(
                {'error': 'domain_id is required'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        # Get mentions with same filtering logic as get_mentions
        show_all = request.GET.get('show_all', 'false').lower() == 'true'
        if show_all:
            mentions = PromptAnalytics.objects.filter(is_published=True, prompt__group__domain_id=domain_id)
        else:
            mentions = PromptAnalytics.objects.filter(is_mention=True, is_published=True, prompt__group__domain_id=domain_id)
        
        # Apply search filter
        search_query = request.GET.get('search', '')
        if search_query:
            mentions = mentions.filter(
                Q(prompt__prompt__icontains=search_query) |
                Q(context_summary__icontains=search_query) |
                Q(prompt__group__domain__name__icontains=search_query)
            )
        
        # Apply platform filter
        platform = request.GET.get('platform', '')
        if platform and platform != 'all' and platform != 'All Platforms':
            mentions = mentions.filter(platform__icontains=platform)
        
        # Apply sentiment filter
        sentiment = request.GET.get('sentiment', '')
        if sentiment and sentiment.lower() != 'all' and sentiment.lower() != 'all sentiments':
            mentions = mentions.filter(sentiment_category__iexact=sentiment)
        
        # Order by most recent first
        mentions = mentions.order_by('-created_at')
        
        # Get count for logging
        mentions_count = mentions.count()
        logger.info(f"Exporting {mentions_count} mentions for domain_id={domain_id}")
        
        # Format mentions data - explicitly evaluate queryset to ensure all results are processed
        # Convert to list first to ensure all results are fetched
        mentions_list = list(mentions)
        logger.info(f"Fetched {len(mentions_list)} mentions from database")
        
        mentions_data = []
        for mention in mentions_list:
            try:
                formatted = _format_mention_for_export(mention)
                mentions_data.append(formatted)
            except Exception as e:
                logger.error(f"Error formatting mention {mention.id}: {str(e)}")
                continue
        
        logger.info(f"Formatted {len(mentions_data)} mentions for export")
        
        # Create Excel file
        excel_file = _create_excel_from_mentions(mentions_data, 'Mentions Export')
        
        # Create HTTP response
        response = HttpResponse(
            excel_file.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="mentions_export_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        
        return response
        
    except Exception as e:
        logger.error(f"Error exporting mentions list: {str(e)}")
        return Response(
            {'error': f'Failed to export mentions: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['GET'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def export_mention_detail(request, analytics_id):
    """
    Export a single mention to Excel file
    """
    try:
        # Get the specific analytics record
        analytics_record = get_object_or_404(PromptAnalytics, id=analytics_id)
        
        # Format the mention data
        mention_data = _format_mention_for_export(analytics_record)
        
        # Create Excel file with single mention
        excel_file = _create_excel_from_mentions([mention_data], f'Mention {analytics_id}')
        
        # Create HTTP response
        response = HttpResponse(
            excel_file.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="mention_{analytics_id}_export_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        
        return response
        
    except Exception as e:
        logger.error(f"Error exporting mention detail: {str(e)}")
        return Response(
            {'error': f'Failed to export mention: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


# ==================== PROMPTS APIs ====================

@api_view(['GET', 'POST'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def prompt_groups_list(request):
    """
    List all prompt groups for the organization or create a new one
    """
    if request.method == 'GET':
        try:
            # Require domain_id and access
            domain_id = request.GET.get('domain_id')
            if not domain_id:
                return Response({'groups': [], 'total_count': 0, 'filters_applied': {'domain_id': None, 'search': request.GET.get('search', '')}})

            user = getattr(request, 'user', None)
            has_access = False
            if user and getattr(user, 'is_authenticated', False):
                if getattr(user, 'role', '') == 'super_admin':
                    has_access = True
                else:
                    has_access = DomainAccess.objects.filter(user=user, domain_id=domain_id).exists()
            if not has_access:
                return Response({'error': 'Forbidden: no access to this domain.'}, status=status.HTTP_403_FORBIDDEN)

            # Get prompt groups scoped to domain
            groups = PromptGroup.objects.filter(domain_id=domain_id)
            
            # Apply search filter if provided
            search_query = request.GET.get('search', '')
            if search_query:
                groups = groups.filter(
                    Q(group_id__icontains=search_query) |
                    Q(domain__name__icontains=search_query)
                )
            
            # Order by creation date
            groups = groups.order_by('-created_at')

            # Pagination
            try:
                limit = int(request.GET.get('limit', 20))
                offset = int(request.GET.get('offset', 0))
            except ValueError:
                limit = 20
                offset = 0
            limit = max(1, min(limit, 100))
            offset = max(0, offset)
            total_count = groups.count()
            groups = groups[offset:offset + limit]
            
            # Prepare response data
            groups_data = []
            for group in groups:
                # Get analytics summary for the group
                analytics = PromptAnalytics.objects.filter(
                    prompt__group=group
                )
                # Derive primary and secondary prompts
                primary_prompt_obj = group.prompts.filter(type='primary').order_by('created_at').first()
                primary_prompt_text = primary_prompt_obj.prompt if primary_prompt_obj else ''
                secondary_prompts_list = list(
                    group.prompts.filter(type='secondary').order_by('created_at').values_list('prompt', flat=True)
                )
                
                # Calculate visibility growth: compare current mentions with previous period (last 7 days vs previous 7 days)
                # Only show growth if there's historical data across multiple time periods
                visibility_growth = None
                
                try:
                    # Get mentions from last 7 days
                    from datetime import timedelta
                    now = timezone.now()
                    seven_days_ago = now - timedelta(days=7)
                    fourteen_days_ago = now - timedelta(days=14)
                    
                    current_period_mentions = analytics.filter(
                        is_mention=True,
                        is_published=True,
                        created_at__gte=seven_days_ago
                    ).count()
                    
                    previous_period_mentions = analytics.filter(
                        is_mention=True,
                        is_published=True,
                        created_at__gte=fourteen_days_ago,
                        created_at__lt=seven_days_ago
                    ).count()
                    
                    # Only calculate growth if we have data in BOTH periods (historical comparison)
                    # This prevents showing 100% growth when there's only data in one period
                    if previous_period_mentions > 0:
                        # We have historical data, calculate growth
                        visibility_growth = round(((current_period_mentions - previous_period_mentions) / previous_period_mentions) * 100, 1)
                    elif current_period_mentions > 0 and previous_period_mentions == 0:
                        # Current period has data but previous doesn't - check if we have ANY older data
                        # If we have data older than 14 days, it means we're tracking but just no data in previous period
                        older_mentions = analytics.filter(
                            is_mention=True,
                            is_published=True,
                            created_at__lt=fourteen_days_ago
                        ).count()
                        
                        if older_mentions > 0:
                            # We have historical data (older than 14 days), so this is a valid comparison
                            # Previous period had 0, current has some = 100% growth
                            visibility_growth = 100.0
                        else:
                            # No historical data at all, don't show growth
                            visibility_growth = None
                    else:
                        # Both periods have 0 mentions, no growth to show
                        visibility_growth = None
                except Exception as e:
                    logger.warning(f"Error calculating visibility growth for group {group.id}: {e}")
                    visibility_growth = None
                
                groups_data.append({
                    'id': group.id,
                    'group_id': group.group_id,
                    'domain_id': group.domain.id,
                    'domain_name': group.domain.name,
                    'total_mentions': group.total_mentions,
                    'total_citations': group.total_citations,
                    'average_position': float(group.average_position),
                    'created_at': group.created_at.isoformat(),
                    'modified_at': group.modified_at.isoformat(),
                    'prompts_count': group.prompts.count(),
                    'primary_prompt': primary_prompt_text,
                    'secondary_prompts': secondary_prompts_list,
                    'visibility_growth': visibility_growth,
                    'track_status': group.track_status,
                    'track_message': group.track_message,
                    'tracked_at': group.tracked_at.isoformat() if group.tracked_at else None,
                    'analytics_summary': {
                        'total_analytics': analytics.count(),
                        'mentions_count': analytics.filter(is_mention=True, is_published=True).count(),
                        'avg_position': float(analytics.aggregate(avg_pos=Avg('position'))['avg_pos'] or 0),
                        'platforms': list(analytics.values_list('platform', flat=True).distinct())
                    }
                })
            
            return Response({
                'groups': groups_data,
                'total_count': total_count,
                'filters_applied': {
                    'domain_id': domain_id,
                    'search': search_query
                },
                'pagination': {
                    'limit': limit,
                    'offset': offset,
                    'returned': len(groups_data)
                }
            })
            
        except Exception as e:
            return Response(
                {'error': f'Failed to retrieve prompt groups: {str(e)}'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )
    
    elif request.method == 'POST':
        try:
            # Validate required fields
            group_id = request.data.get('group_id')
            domain_id = request.data.get('domain_id')
            primary_prompts = request.data.get('primary_prompts', [])
            secondary_prompts = request.data.get('secondary_prompts', [])
            
            if not group_id or not domain_id:
                return Response(
                    {'error': 'group_id and domain_id are required'},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            if not primary_prompts and not secondary_prompts:
                return Response(
                    {'error': 'At least one primary or secondary prompt is required'},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            # Check if group already exists
            if PromptGroup.objects.filter(group_id=group_id, domain_id=domain_id).exists():
                return Response(
                    {'error': 'Prompt group with this ID already exists for this domain'},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            # Get domain
            domain = get_object_or_404(Domain, id=domain_id)
            
            # Create new prompt group
            group = PromptGroup.objects.create(
                group_id=group_id,
                domain=domain
            )
            
            # Platforms to create analytics for
            platforms = ["ChatGPT", "Google Gemini", "Perplexity"]
            
            # Create prompts and analytics
            created_prompts = []
            created_analytics = []
            
            # Create primary prompts
            for prompt_text in primary_prompts:
                if prompt_text.strip():
                    prompt = Prompt.objects.create(
                        prompt=prompt_text.strip(),
                        group=group,
                        track_status='INIT',
                        type='primary'
                    )
                    created_prompts.append(prompt)
                    
                    # Create analytics for each platform
                    for platform in platforms:
                        analytics = PromptAnalytics.objects.create(
                            prompt=prompt,
                            platform=platform,
                            is_mention=False,  # Initially not a mention
                            position=0.0,
                            sentiment_category='neutral',
                            sentiment_score=0.0,
                            total_mentions=0,
                            total_citations=0
                        )
                        created_analytics.append(analytics)
            
            # Create secondary prompts
            for prompt_text in secondary_prompts:
                if prompt_text.strip():
                    prompt = Prompt.objects.create(
                        prompt=prompt_text.strip(),
                        group=group,
                        track_status='INIT',
                        type='secondary'
                    )
                    created_prompts.append(prompt)
                    
                    # Create analytics for each platform
                    for platform in platforms:
                        analytics = PromptAnalytics.objects.create(
                            prompt=prompt,
                            platform=platform,
                            is_mention=False,  # Initially not a mention
                            position=0.0,
                            sentiment_category='neutral',
                            sentiment_score=0.0,
                            total_mentions=0,
                            total_citations=0
                        )
                        created_analytics.append(analytics)
            
            return Response({
                'message': 'Prompt group created successfully',
                'group': {
                    'id': group.id,
                    'group_id': group.group_id,
                    'domain_id': group.domain.id,
                    'domain_name': group.domain.name,
                    'created_at': group.created_at.isoformat()
                },
                'prompts_created': len(created_prompts),
                'analytics_created': len(created_analytics),
                'platforms': platforms
            }, status=status.HTTP_201_CREATED)
            
        except Exception as e:
            return Response(
                {'error': f'Failed to create prompt group: {str(e)}'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )


@api_view(['GET', 'PUT', 'DELETE'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def prompt_group_detail(request, group_id):
    """
    Get, update, or delete a specific prompt group
    Optional query parameter: platform - filter metrics by platform
    """
    try:
        group = get_object_or_404(PromptGroup, id=group_id)
        
        if request.method == 'GET':
            # Get platform filter from query parameter
            platform_filter = request.GET.get('platform')
            
            # Get detailed information about the group
            # Use prefetch_related to avoid N+1 queries when accessing prompt analytics
            prompts = group.prompts.prefetch_related(
                'analytics'  # Prefetch all analytics for each prompt
            ).all().order_by('created_at')

            # Get all analytics for the group (for summary stats)
            analytics = PromptAnalytics.objects.filter(prompt__group=group)
            # Analytics filtering is done per-prompt below to show all variants with platform-specific metrics

            # Derive primary and secondary prompts for detail view
            primary_prompt_obj = group.prompts.filter(type='primary').order_by('created_at').first()
            primary_prompt_text = primary_prompt_obj.prompt if primary_prompt_obj else ''
            secondary_prompts_list = list(
                group.prompts.filter(type='secondary').order_by('created_at').values_list('prompt', flat=True)
            )
            
            # Get full_ai_response from latest analytics
            # If platform_filter is provided, get response for that specific platform
            # Otherwise, get response from primary prompt
            primary_full_ai_response = ''

            # Build the base query - if platform_filter exists, search across all prompts in group for that platform
            # Otherwise, only search within primary_prompt_obj
            if platform_filter:
                # When filtering by platform, get the response from ANY prompt in the group that matches the platform
                base_query = PromptAnalytics.objects.filter(
                    prompt__group=group,
                    platform=platform_filter
                )
            elif primary_prompt_obj:
                # No platform filter, get from primary prompt
                base_query = PromptAnalytics.objects.filter(
                    prompt=primary_prompt_obj
                )
            else:
                base_query = PromptAnalytics.objects.none()

            # Try multiple fallbacks to find the best available response
            if base_query.exists():
                # First try: Get latest published mention analytics
                latest_primary_analytic = base_query.filter(
                    is_mention=True,
                    is_published=True
                ).order_by('-created_at').first()

                # Fallback 1: If no mention found, try any published analytics
                if not latest_primary_analytic:
                    latest_primary_analytic = base_query.filter(
                        is_published=True
                    ).order_by('-created_at').first()

                # Fallback 2: If still nothing, try any analytics record (even if not published)
                if not latest_primary_analytic:
                    latest_primary_analytic = base_query.order_by('-created_at').first()

                if latest_primary_analytic and latest_primary_analytic.context_summary:
                    primary_full_ai_response = latest_primary_analytic.context_summary

            prompts_data = []
            for prompt in prompts:
                # Use prefetched analytics instead of making new queries
                all_prompt_analytics = prompt.analytics.all()

                # Get platform-filtered analytics for metrics calculation
                if platform_filter:
                    # Filter in Python (analytics already prefetched)
                    prompt_analytics = [a for a in all_prompt_analytics if a.platform == platform_filter]
                else:
                    # Use all analytics for this prompt
                    prompt_analytics = list(all_prompt_analytics)
                
                # Sum total_mentions from analytics records, not just count records
                # This matches how group.total_mentions is calculated (Sum of total_mentions field)
                mention_analytics = [a for a in prompt_analytics if a.is_mention and a.is_published]
                prompt_mentions = sum(analytic.total_mentions or 1 for analytic in mention_analytics)
                
                # Calculate citations count
                total_citations = 0
                for analytic in prompt_analytics:
                    citation_list = getattr(analytic, 'citation_list', []) or []
                    total_citations += len(citation_list) if citation_list else 0
                
                # Calculate sentiment breakdown - use sentiment_category directly (like Mentions page)
                # This matches how mentions display sentiment, using the actual sentiment_category field
                total_positive_mentions = 0
                total_neutral_mentions = 0
                total_negative_mentions = 0
                
                # Only count mentions (is_mention=True, is_published=True) for sentiment calculation
                for analytic in mention_analytics:  # Already filtered above
                    # Use sentiment_category directly (like Mentions page does)
                    sentiment_category = getattr(analytic, 'sentiment_category', None)
                    # Count each mention (not just analytics record)
                    mentions = analytic.total_mentions or 1
                    
                    # Normalize sentiment_category to lowercase for comparison
                    if sentiment_category:
                        sentiment_lower = sentiment_category.lower()
                        if sentiment_lower == 'positive':
                            total_positive_mentions += mentions
                        elif sentiment_lower == 'negative':
                            total_negative_mentions += mentions
                        else:
                            # Default to neutral if not positive or negative
                            total_neutral_mentions += mentions
                    else:
                        # If no sentiment_category, try to derive from sentiment_score as fallback
                        sentiment_score = float(analytic.sentiment_score or 0)
                        if sentiment_score > 0.33:
                            total_positive_mentions += mentions
                        elif sentiment_score < -0.33:
                            total_negative_mentions += mentions
                        else:
                            total_neutral_mentions += mentions
                
                total_sentiment_mentions = total_positive_mentions + total_neutral_mentions + total_negative_mentions
                sentiment_percentages = {
                    'positive': round((total_positive_mentions / total_sentiment_mentions * 100) if total_sentiment_mentions > 0 else 0, 1),
                    'neutral': round((total_neutral_mentions / total_sentiment_mentions * 100) if total_sentiment_mentions > 0 else 0, 1),
                    'negative': round((total_negative_mentions / total_sentiment_mentions * 100) if total_sentiment_mentions > 0 else 0, 1)
                }
                
                # Get latest analytics for position (most recent tracked date)
                latest_analytic = max(prompt_analytics, key=lambda a: a.created_at) if prompt_analytics else None
                latest_position = float(latest_analytic.position) if latest_analytic else 0

                # Get latest mention analytics ID for navigation
                mention_analytics_sorted = sorted(
                    mention_analytics,
                    key=lambda a: a.created_at,
                    reverse=True
                )
                latest_mention_id = mention_analytics_sorted[0].id if mention_analytics_sorted else None

                # Get all platforms that have analytics for this prompt (for filtering)
                # Use all_prompt_analytics to get complete platform list, not filtered
                all_platforms = list(set(a.platform for a in all_prompt_analytics if a.platform))
                # Ensure we have at least the latest platform if available
                latest_all_analytic = max(all_prompt_analytics, key=lambda a: a.created_at) if all_prompt_analytics else None
                if latest_all_analytic and latest_all_analytic.platform and latest_all_analytic.platform not in all_platforms:
                    all_platforms.append(latest_all_analytic.platform)
                
                prompts_data.append({
                    'id': prompt.id,
                    'prompt_text': prompt.prompt,
                    'track_status': prompt.track_status,
                    'type': prompt.type,
                    'last_tracked_at': prompt.tracked_at.isoformat() if prompt.tracked_at else None,
                    'track_message': prompt.track_message,
                    'created_at': prompt.created_at.isoformat(),
                    'analytics_count': len(prompt_analytics),
                    'mentions_count': prompt_mentions,
                    'avg_position': sum(a.position for a in prompt_analytics if a.position) / len([a for a in prompt_analytics if a.position]) if any(a.position for a in prompt_analytics) else 0.0,
                    'latest_position': latest_position,
                    'citations_count': total_citations,
                    'sentiment': sentiment_percentages,
                    'platform': latest_analytic.platform if latest_analytic else None,
                    'platforms': all_platforms,  # All platforms for this prompt
                    'latest_mention_id': latest_mention_id
                })

            # Platform distribution for the group - use PromptGroupMetricSnapshot
            # Get the latest snapshots for this group (prefer daily, fallback to weekly/monthly)
            now = timezone.now().date()
            # Get snapshots from the last 30 days to get recent data
            start_date = now - timedelta(days=30)
            
            # Get platform-specific snapshots (exclude NULL/empty platforms)
            platform_snapshots = PromptGroupMetricSnapshot.objects.filter(
                prompt_group=group,
                snapshot_date__gte=start_date,
                snapshot_date__lte=now
            ).exclude(platform__isnull=True).exclude(platform='')
            
            # Aggregate by platform (sum mentions, weighted average for position)
            platform_aggregates = {}
            for snapshot in platform_snapshots:
                platform = snapshot.platform
                if platform not in platform_aggregates:
                    platform_aggregates[platform] = {
                        'mention_count': 0,
                        'position_sum': 0,
                        'position_weight': 0
                    }
                
                platform_aggregates[platform]['mention_count'] += snapshot.mentions
                if snapshot.average_position and snapshot.mentions > 0:
                    platform_aggregates[platform]['position_sum'] += float(snapshot.average_position) * snapshot.mentions
                    platform_aggregates[platform]['position_weight'] += snapshot.mentions
            
            # Build platform distribution list with weighted averages
            platform_dist = []
            for platform, agg in platform_aggregates.items():
                avg_pos = (agg['position_sum'] / agg['position_weight']) if agg['position_weight'] > 0 else 0
                platform_dist.append({
                    'platform': platform,
                    'count': int(agg['mention_count']),
                    'avg_position': float(avg_pos)
                })

            # Fallback: if no snapshot data, calculate from analytics directly
            if not platform_dist:
                # Get all analytics for the group
                all_analytics = PromptAnalytics.objects.filter(
                    prompt__group=group,
                    is_mention=True,
                    is_published=True
                ).exclude(platform__isnull=True).exclude(platform='')

                # Aggregate by platform
                from django.db.models import Count, Avg
                platform_stats = all_analytics.values('platform').annotate(
                    count=Count('id'),
                    avg_position=Avg('position')
                ).order_by('-count')

                for stat in platform_stats:
                    platform_dist.append({
                        'platform': stat['platform'],
                        'count': stat['count'],
                        'avg_position': float(stat['avg_position'] or 0)
                    })

            # Sort by count descending
            platform_dist.sort(key=lambda x: x['count'], reverse=True)

            # Variants performance based on mentions per prompt - use PromptMetricSnapshot
            variants_perf = []
            now_date = timezone.now().date()
            start_date_snapshots = now_date - timedelta(days=30)  # Last 30 days
            
            for p in prompts:
                # Get snapshots for this prompt (aggregated across all platforms)
                prompt_snapshots = PromptMetricSnapshot.objects.filter(
                    prompt=p,
                    snapshot_date__gte=start_date_snapshots,
                    snapshot_date__lte=now_date
                )
                # Sum mentions from all snapshots
                total_mentions = prompt_snapshots.aggregate(total=Sum('mentions'))['total'] or 0
                variants_perf.append({
                    'prompt_id': p.id,
                    'prompt_text': p.prompt,
                    'mentions': int(total_mentions)
                })

            # Mention trends (by month in last 6 months) - use PromptGroupMetricSnapshot
            from django.utils import timezone as _tz
            from datetime import timedelta as _td
            end_date = _tz.now().date()
            start_date = end_date - _td(days=180)
            trends = []
            
            # Group snapshots by month
            group_snapshots = PromptGroupMetricSnapshot.objects.filter(
                prompt_group=group,
                snapshot_date__gte=start_date,
                snapshot_date__lte=end_date
            )
            
            # Aggregate by month
            monthly_data = {}
            for snapshot in group_snapshots:
                month_key = snapshot.snapshot_date.strftime('%Y-%m')
                if month_key not in monthly_data:
                    monthly_data[month_key] = {
                        'mentions': 0,
                        'position_sum': 0,
                        'position_weight': 0
                    }
                
                monthly_data[month_key]['mentions'] += snapshot.mentions
                if snapshot.average_position and snapshot.mentions > 0:
                    monthly_data[month_key]['position_sum'] += float(snapshot.average_position) * snapshot.mentions
                    monthly_data[month_key]['position_weight'] += snapshot.mentions
            
            # Generate trends for each month in the range
            cur = start_date.replace(day=1)  # Start from first day of month
            while cur <= end_date:
                month_key = cur.strftime('%Y-%m')
                month_data = monthly_data.get(month_key, {'mentions': 0, 'position_sum': 0, 'position_weight': 0})
                avg_pos = (month_data['position_sum'] / month_data['position_weight']) if month_data['position_weight'] > 0 else 0
                
                trends.append({
                    'date': month_key,
                    'mentions': int(month_data['mentions']),
                    'avg_position': float(avg_pos)
                })
                
                # Move to next month
                if cur.month == 12:
                    cur = cur.replace(year=cur.year+1, month=1, day=1)
                else:
                    cur = cur.replace(month=cur.month+1, day=1)
            
            # Calculate visibility growth: compare current mentions with previous period (last 7 days vs previous 7 days)
            # Use PromptGroupMetricSnapshot
            visibility_growth = None
            try:
                now_date = _tz.now().date()
                seven_days_ago = now_date - _td(days=7)
                fourteen_days_ago = now_date - _td(days=14)
                
                # Get current period snapshots (last 7 days)
                current_period_snapshots = PromptGroupMetricSnapshot.objects.filter(
                    prompt_group=group,
                    snapshot_date__gte=seven_days_ago,
                    snapshot_date__lte=now_date
                )
                current_period_mentions = current_period_snapshots.aggregate(total=Sum('mentions'))['total'] or 0
                
                # Get previous period snapshots (7-14 days ago)
                previous_period_snapshots = PromptGroupMetricSnapshot.objects.filter(
                    prompt_group=group,
                    snapshot_date__gte=fourteen_days_ago,
                    snapshot_date__lt=seven_days_ago
                )
                previous_period_mentions = previous_period_snapshots.aggregate(total=Sum('mentions'))['total'] or 0
                
                # Calculate percentage growth
                if previous_period_mentions > 0:
                    visibility_growth = round(((current_period_mentions - previous_period_mentions) / previous_period_mentions) * 100, 1)
                elif current_period_mentions > 0:
                    # If previous period had 0 mentions but current has some, show 100% growth
                    visibility_growth = 100.0
                else:
                    # Both periods have 0 mentions, no growth
                    visibility_growth = 0.0
            except Exception as e:
                logger.warning(f"Error calculating visibility growth for group {group.id}: {e}")
                visibility_growth = None
            
            # Calculate average_position from PromptGroupMetricSnapshot (same as graph) for consistency
            # Use the same snapshots used for trends calculation
            calculated_avg_position = 0.0
            try:
                # Get all snapshots used in trends (last 180 days)
                all_snapshots = PromptGroupMetricSnapshot.objects.filter(
                    prompt_group=group,
                    snapshot_date__gte=start_date,
                    snapshot_date__lte=end_date
                )
                
                position_sum = 0
                position_weight = 0
                for snapshot in all_snapshots:
                    if snapshot.average_position and snapshot.mentions > 0:
                        position_sum += float(snapshot.average_position) * snapshot.mentions
                        position_weight += snapshot.mentions
                
                if position_weight > 0:
                    calculated_avg_position = float(position_sum / position_weight)
                else:
                    # Fallback to model field if no snapshot data
                    calculated_avg_position = float(group.average_position)
            except Exception as e:
                logger.warning(f"Error calculating average position for group {group.id}: {e}")
                calculated_avg_position = float(group.average_position)

            return Response({
                'group': {
                    'id': group.id,
                    'group_id': group.group_id,
                    'domain_id': group.domain.id,
                    'domain_name': group.domain.name,
                    'theme': group.theme or '',
                    'primary_prompt': primary_prompt_text,
                    'primary_full_ai_response': primary_full_ai_response,
                    'secondary_prompts': secondary_prompts_list,
                    'total_mentions': group.total_mentions,
                    'total_citations': group.total_citations,
                    'average_position': calculated_avg_position,
                    'visibility_growth': visibility_growth,
                    'created_at': group.created_at.isoformat(),
                    'modified_at': group.modified_at.isoformat(),
                    'prompts': prompts_data,
                    'active_variants': prompts.count(),
                    'analytics_summary': {
                        'total_analytics': analytics.count(),
                        'mentions_count': analytics.filter(is_mention=True, is_published=True).count(),
                        'platforms': list(analytics.values_list('platform', flat=True).distinct()),
                        'sentiments': list(analytics.values_list('sentiment_category', flat=True).distinct())
                    },
                    'platform_distribution': platform_dist,
                    'variants_performance': variants_perf,
                    'mention_trends': trends
                }
            })
        
        elif request.method == 'PUT':
            # Update group information and prompts (primary + variants)
            group_id_new = request.data.get('group_id', group.group_id)
            domain_id = request.data.get('domain_id', group.domain.id)
            primary_prompt_text = request.data.get('primary_prompt')
            secondary_prompts_texts = request.data.get('secondary_prompts', []) or []

            # Check if new group_id already exists (if changed)
            if group_id_new != group.group_id:
                if PromptGroup.objects.filter(group_id=group_id_new, domain_id=domain_id).exists():
                    return Response(
                        {'error': 'Prompt group with this ID already exists for this domain'},
                        status=status.HTTP_400_BAD_REQUEST
                    )

            # Update group fields
            group.group_id = group_id_new
            group.domain_id = domain_id
            group.save()

            # If primary/secondary provided, reconcile prompts to match
            if primary_prompt_text is not None:
                primary_prompt_text = str(primary_prompt_text).strip()
                # Normalize secondary list (unique, trimmed, exclude empty and primary)
                normalized_secondaries = []
                for s in secondary_prompts_texts:
                    s_norm = str(s).strip()
                    if s_norm and s_norm != primary_prompt_text and s_norm not in normalized_secondaries:
                        normalized_secondaries.append(s_norm)

                # Ensure primary prompt exists and is marked primary
                primary_prompt_obj = group.prompts.filter(prompt=primary_prompt_text).first()
                if primary_prompt_obj:
                    if primary_prompt_obj.type != 'primary':
                        primary_prompt_obj.type = 'primary'
                        primary_prompt_obj.save()
                else:
                    primary_prompt_obj = Prompt.objects.create(
                        prompt=primary_prompt_text,
                        group=group,
                        track_status='INIT',
                        type='primary'
                    )

                # Upsert secondary prompts
                existing_prompts = {p.prompt: p for p in group.prompts.all()}
                for sec_text in normalized_secondaries:
                    if sec_text in existing_prompts:
                        p = existing_prompts[sec_text]
                        if p.type != 'secondary':
                            p.type = 'secondary'
                            p.save()
                    else:
                        Prompt.objects.create(
                            prompt=sec_text,
                            group=group,
                            track_status='INIT',
                            type='secondary'
                        )

                # Remove any prompts no longer in the provided set (except keep analytics integrity)
                keep_set = set([primary_prompt_text] + normalized_secondaries)
                for p in group.prompts.exclude(prompt__in=keep_set):
                    p.delete()

                # Ensure only one primary remains
                group.prompts.exclude(id=primary_prompt_obj.id).filter(type='primary').update(type='secondary')

            return Response({
                'message': 'Prompt group updated successfully',
                'group': {
                    'id': group.id,
                    'group_id': group.group_id,
                    'domain_id': group.domain.id,
                    'domain_name': group.domain.name,
                    'modified_at': group.modified_at.isoformat()
                }
            })
        
        elif request.method == 'DELETE':
            # Delete the group (this will cascade delete all prompts and analytics)
            group.delete()
            return Response({
                'message': 'Prompt group deleted successfully'
            })
            
    except Exception as e:
        return Response(
            {'error': f'Failed to process prompt group: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['GET', 'POST'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def prompts_list(request):
    """
    List all prompts for the organization or create a new one
    """
    if request.method == 'GET':
        try:
            # Require domain_id and access
            domain_id = request.GET.get('domain_id')
            if not domain_id:
                return Response({'prompts': [], 'total_count': 0, 'filters_applied': {
                    'group_id': request.GET.get('group_id'),
                    'domain_id': None,
                    'track_status': request.GET.get('track_status'),
                    'type': request.GET.get('type'),
                    'search': request.GET.get('search', '')
                }})

            user = getattr(request, 'user', None)
            has_access = False
            if user and getattr(user, 'is_authenticated', False):
                if getattr(user, 'role', '') == 'super_admin':
                    has_access = True
                else:
                    has_access = DomainAccess.objects.filter(user=user, domain_id=domain_id).exists()
            if not has_access:
                return Response({'error': 'Forbidden: no access to this domain.'}, status=status.HTTP_403_FORBIDDEN)

            # Get prompts scoped to domain via group
            prompts = Prompt.objects.filter(group__domain_id=domain_id)
            
            # Apply filters
            group_id = request.GET.get('group_id')
            if group_id:
                prompts = prompts.filter(group_id=group_id)
            
            track_status = request.GET.get('track_status')
            if track_status:
                prompts = prompts.filter(track_status=track_status)
            
            prompt_type = request.GET.get('type')
            if prompt_type:
                prompts = prompts.filter(type=prompt_type)
            
            # Apply search filter
            search_query = request.GET.get('search', '')
            if search_query:
                prompts = prompts.filter(prompt__icontains=search_query)
            
            # Order by creation date
            prompts = prompts.order_by('-created_at')

            # Pagination
            try:
                limit = int(request.GET.get('limit', 20))
                offset = int(request.GET.get('offset', 0))
            except ValueError:
                limit = 20
                offset = 0
            limit = max(1, min(limit, 100))
            offset = max(0, offset)
            total_count = prompts.count()
            prompts = prompts[offset:offset + limit]
            
            # Prepare response data
            prompts_data = []
            for prompt in prompts:
                # Get analytics for this prompt
                analytics = PromptAnalytics.objects.filter(
                    prompt=prompt
                )
                
                prompts_data.append({
                    'id': prompt.id,
                    'prompt_text': prompt.prompt,
                    'group_id': prompt.group.group_id if prompt.group else None,
                    'group_name': f"Group {prompt.group.group_id}" if prompt.group else "No Group",
                    'domain_id': prompt.group.domain.id if prompt.group else None,
                    'domain_name': prompt.group.domain.name if prompt.group else None,
                    'track_status': prompt.track_status,
                    'type': prompt.type,
                    'last_tracked_at': prompt.tracked_at.isoformat() if prompt.tracked_at else None,
                    'track_message': prompt.track_message,
                    'created_at': prompt.created_at.isoformat(),
                    'modified_at': prompt.modified_at.isoformat(),
                    'analytics_summary': {
                        'total_analytics': analytics.count(),
                        'mentions_count': analytics.filter(is_mention=True, is_published=True).count(),
                        'avg_position': float(analytics.aggregate(avg_pos=Avg('position'))['avg_pos'] or 0),
                        'platforms': list(analytics.values_list('platform', flat=True).distinct()),
                        'latest_mention': analytics.filter(is_mention=True, is_published=True).order_by('-created_at').first().created_at.isoformat() if analytics.filter(is_mention=True, is_published=True).exists() else None
                    }
                })
            
            return Response({
                'prompts': prompts_data,
                'total_count': total_count,
                'filters_applied': {
                    'group_id': group_id,
                    'domain_id': domain_id,
                    'track_status': track_status,
                    'type': prompt_type,
                    'search': search_query
                },
                'pagination': {
                    'limit': limit,
                    'offset': offset,
                    'returned': len(prompts_data)
                }
            })
            
        except Exception as e:
            return Response(
                {'error': f'Failed to retrieve prompts: {str(e)}'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )
    
    elif request.method == 'POST':
        try:
            # Validate required fields
            prompt_text = request.data.get('prompt')
            group_id = request.data.get('group_id')
            
            if not prompt_text or not group_id:
                return Response(
                    {'error': 'prompt and group_id are required'},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            # Optional fields
            track_status = request.data.get('track_status', 'INIT')
            prompt_type = request.data.get('type', 'primary')
            track_message = request.data.get('track_message', '')
            
            # Create new prompt (group required)
            group = get_object_or_404(PromptGroup, id=group_id)
            prompt = Prompt.objects.create(
                prompt=prompt_text,
                group=group,
                track_status=track_status,
                type=prompt_type,
                track_message=track_message
            )
            
            return Response({
                'message': 'Prompt created successfully',
                'prompt': {
                    'id': prompt.id,
                    'prompt_text': prompt.prompt,
                    'group_id': prompt.group.group_id if prompt.group else None,
                    'domain_id': prompt.group.domain.id if prompt.group else None,
                    'domain_name': prompt.group.domain.name if prompt.group else None,
                    'track_status': prompt.track_status,
                    'type': prompt.type,
                    'created_at': prompt.created_at.isoformat()
                }
            }, status=status.HTTP_201_CREATED)
            
        except Exception as e:
            return Response(
                {'error': f'Failed to create prompt: {str(e)}'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )


@api_view(['GET', 'PUT', 'DELETE'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def prompt_detail(request, prompt_id):
    """
    Get, update, or delete a specific prompt
    """
    try:
        prompt = get_object_or_404(Prompt, id=prompt_id)
        
        if request.method == 'GET':
            # Get detailed information about the prompt
            analytics = PromptAnalytics.objects.filter(
                prompt=prompt
            ).order_by('-created_at')
            
            analytics_data = []
            for analytic in analytics:
                analytics_data.append({
                    'id': analytic.id,
                    'platform': analytic.platform,
                    'is_mention': analytic.is_mention,
                    'total_mentions': analytic.total_mentions,
                    'total_citations': analytic.total_citations,
                    'position': float(analytic.position),
                    'sentiment': getattr(analytic, 'sentiment_category', 'neutral'),
                    'sentiment_score': float(analytic.sentiment_score),
                    'context_summary': analytic.context_summary,
                    'views': analytic.views,
                    'shares': analytic.shares,
                    'engagement_score': float(analytic.engagement_score),
                    'created_at': analytic.created_at.isoformat(),
                    'citations_count': len(getattr(analytic, 'citation_list', []) or [])
                })
            
            return Response({
                'prompt': {
                    'id': prompt.id,
                    'prompt_text': prompt.prompt,
                    'group_id': prompt.group.group_id if prompt.group else None,
                    'group_name': f"Group {prompt.group.group_id}" if prompt.group else "No Group",
                    'domain_id': prompt.group.domain.id if prompt.group else None,
                    'domain_name': prompt.group.domain.name if prompt.group else None,
                    'track_status': prompt.track_status,
                    'type': prompt.type,
                    'last_tracked_at': prompt.tracked_at.isoformat() if prompt.tracked_at else None,
                    'track_message': prompt.track_message,
                    'created_at': prompt.created_at.isoformat(),
                    'modified_at': prompt.modified_at.isoformat(),
                    'analytics': analytics_data,
                    'analytics_summary': {
                        'total_analytics': analytics.count(),
                        'mentions_count': analytics.filter(is_mention=True, is_published=True).count(),
                        'platforms': list(analytics.values_list('platform', flat=True).distinct()),
                        'avg_position': float(analytics.aggregate(avg_pos=Avg('position'))['avg_pos'] or 0),
                        'avg_sentiment': float(analytics.aggregate(avg_sent=Avg('sentiment_score'))['avg_sent'] or 0)
                    }
                }
            })
        
        elif request.method == 'PUT':
            # Update prompt information
            prompt.prompt = request.data.get('prompt', prompt.prompt)
            prompt.group_id = request.data.get('group_id', prompt.group_id)
            prompt.track_status = request.data.get('track_status', prompt.track_status)
            prompt.type = request.data.get('type', prompt.type)
            prompt.track_message = request.data.get('track_message', prompt.track_message)
            prompt.save()
            
            return Response({
                'message': 'Prompt updated successfully',
                'prompt': {
                    'id': prompt.id,
                    'prompt_text': prompt.prompt,
                    'group_id': prompt.group.group_id if prompt.group else None,
                    'track_status': prompt.track_status,
                    'type': prompt.type,
                    'modified_at': prompt.modified_at.isoformat()
                }
            })
        
        elif request.method == 'DELETE':
            # Delete the prompt (this will cascade delete all analytics)
            prompt.delete()
            return Response({
                'message': 'Prompt deleted successfully'
            })
            
    except Exception as e:
        return Response(
            {'error': f'Failed to process prompt: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def bulk_update_prompts(request):
    """
    Bulk update prompts (track status, type, etc.)
    """
    try:
        prompt_ids = request.data.get('prompt_ids', [])
        updates = request.data.get('updates', {})
        
        if not prompt_ids:
            return Response(
                {'error': 'prompt_ids is required'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        # Validate updates
        allowed_updates = ['track_status', 'type', 'track_message', 'group_id']
        for key in updates.keys():
            if key not in allowed_updates:
                return Response(
                    {'error': f'Invalid update field: {key}'},
                    status=status.HTTP_400_BAD_REQUEST
                )
        
        # Update prompts
        updated_count = 0
        for prompt_id in prompt_ids:
            try:
                prompt = Prompt.objects.get(id=prompt_id)
                for key, value in updates.items():
                    setattr(prompt, key, value)
                prompt.save()
                updated_count += 1
            except Prompt.DoesNotExist:
                continue
        
        return Response({
            'message': f'Successfully updated {updated_count} prompts',
            'updated_count': updated_count,
            'total_requested': len(prompt_ids)
        })
        
    except Exception as e:
        return Response(
            {'error': f'Failed to bulk update prompts: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['GET'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def prompt_analytics(request, prompt_id):
    """
    Get analytics for a specific prompt
    """
    try:
        prompt = get_object_or_404(Prompt, id=prompt_id)
        
        # Get analytics for this prompt
        analytics = PromptAnalytics.objects.filter(
            prompt=prompt
        ).order_by('-created_at')
        
        # Apply filters
        platform = request.GET.get('platform')
        if platform:
            analytics = analytics.filter(platform=platform)
        
        is_mention = request.GET.get('is_mention')
        if is_mention is not None:
            analytics = analytics.filter(is_mention=is_mention.lower() == 'true')
        
        # Prepare response data
        analytics_data = []
        for analytic in analytics:
            citations_data = []
            citations_src = getattr(analytic, 'citation_list', None) or getattr(analytic, 'citations', []) or []
            if citations_src:
                for i, citation in enumerate(citations_src):
                    citations_data.append({
                        'id': i + 1,
                        'text': citation.get('text', ''),
                        'source': citation.get('source', ''),
                        'url': citation.get('url', ''),
                        'description': citation.get('description', '')
                    })
            
            analytics_data.append({
                'id': analytic.id,
                'platform': analytic.platform,
                'is_mention': analytic.is_mention,
                'total_mentions': analytic.total_mentions,
                'total_citations': analytic.total_citations,
                'position': float(analytic.position),
                'sentiment': getattr(analytic, 'sentiment_category', 'neutral'),
                'sentiment_score': float(analytic.sentiment_score),
                'context_summary': analytic.context_summary,
                'views': analytic.views,
                'shares': analytic.shares,
                'engagement_score': float(analytic.engagement_score),
                'competitor_mentions': (getattr(analytic, 'competitor_mention_list', None) or getattr(analytic, 'competitor_mentions', []) or []),
                'key_topics': (getattr(analytic, 'topic_list', None) or getattr(analytic, 'key_topics', []) or []),
                'created_at': analytic.created_at.isoformat(),
                'citations': citations_data,
                'citations_count': len(citations_data)
            })
        
        return Response({
            'prompt': {
                'id': prompt.id,
                'prompt_text': prompt.prompt,
                'group_id': prompt.group.group_id if prompt.group else None,
                'domain_name': (prompt.group.domain.name if prompt.group else None)
            },
            'analytics': analytics_data,
            'total_count': len(analytics_data),
            'filters_applied': {
                'platform': platform,
                'is_mention': is_mention
            },
            'summary': {
                'total_analytics': analytics.count(),
                'mentions_count': analytics.filter(is_mention=True, is_published=True).count(),
                'platforms': list(analytics.values_list('platform', flat=True).distinct()),
                'avg_position': float(analytics.aggregate(avg_pos=Avg('position'))['avg_pos'] or 0),
                'avg_sentiment': float(analytics.aggregate(avg_sent=Avg('sentiment_score'))['avg_sent'] or 0)
            }
        })
        
    except Exception as e:
        return Response(
            {'error': f'Failed to retrieve prompt analytics: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['GET'])
@permission_classes([AllowAny])
def get_historical_trends(request):
    """
    Get comprehensive historical trends data for the Historical Trends page
    Uses DomainMetricSnapshot for efficient time-series queries
    Returns: visibility progression, platform growth, competitor comparison, seasonal patterns, summary metrics
    """
    try:
        domain_id = request.GET.get('domain_id')
        months = int(request.GET.get('months', 12))
        start_date_qs = request.GET.get('start_date')  # YYYY-MM-DD (optional)
        end_date_qs = request.GET.get('end_date')      # YYYY-MM-DD (optional)
        
        if not domain_id:
            return Response({'error': 'domain_id is required'}, status=status.HTTP_400_BAD_REQUEST)
        
        # Calculate date range
        from datetime import date as date_class
        today = date_class.today()
        end_date = today
        if end_date_qs:
            try:
                end_date = datetime.strptime(end_date_qs, '%Y-%m-%d').date()
            except Exception:
                pass
        if start_date_qs:
            try:
                start_date = datetime.strptime(start_date_qs, '%Y-%m-%d').date()
            except Exception:
                start_date = end_date - timedelta(days=months * 30)
        else:
            start_date = end_date - timedelta(days=months * 30)
        
        # Use monthly snapshots if available, fallback to weekly/daily
        # Try monthly first, then weekly, then daily
        # Note: DomainMetricSnapshot records are created per platform (not aggregated)
        # So we need to query platform-specific snapshots and aggregate them
        period_types = ['monthly', 'weekly', 'daily']
        snapshots = None
        selected_period_type = None
        
        for period_type in period_types:
            snapshots = DomainMetricSnapshot.objects.filter(
                domain_id=domain_id,
                snapshot_date__gte=start_date,
                snapshot_date__lte=end_date,
                period_type=period_type
            ).exclude(platform__isnull=True).exclude(platform='')  # Only platform-specific snapshots
            
            if snapshots.exists():
                selected_period_type = period_type
                break
        
        # If no snapshots found, return empty data
        if not snapshots or not snapshots.exists():
            return Response({
                'visibility_trend': [],
                'platform_growth': [],
                'competitor_comparison': [],
                'seasonal_pattern': [],
                'summary': {
                    'visibility_growth': 0,
                    'mention_growth': 0,
                    'position_improvement': 0,
                    'market_share_gain': 0
                },
                'period_months': months,
                'start_date': start_date.isoformat(),
                'end_date': end_date.isoformat()
            })
        
        # Helper to get month key from date
        def get_month_key(d):
            return d.strftime('%Y-%m')
        
        def get_month_name(d):
            return d.strftime('%b')
        
        # Map platform names from database to frontend keys
        platform_mapping = {
            'ChatGPT': 'chatgpt',
            'Google Gemini': 'gemini',
            'Gemini': 'gemini',
            'Perplexity': 'perplexity',
            'Claude': 'claude'
        }
        
        # Aggregate snapshots by month (across all platforms for overall metrics)
        monthly_data = {}
        for snapshot in snapshots:
            month_key = get_month_key(snapshot.snapshot_date)
            if month_key not in monthly_data:
                monthly_data[month_key] = {
                    'month': get_month_name(snapshot.snapshot_date),
                    'mentions': 0,
                    'avg_position': 0.0,
                    'sentiment': 0.0,
                    'visibility_score': 0.0,
                    'position_sum': 0.0,
                    'sentiment_sum': 0.0,
                    'visibility_sum': 0.0,
                    'total_mentions': 0,
                    'platforms': {}
                }
            
            data = monthly_data[month_key]
            # Aggregate overall metrics (weighted by mentions)
            data['mentions'] += snapshot.mentions
            data['total_mentions'] += snapshot.mentions
            data['position_sum'] += float(snapshot.average_position) * snapshot.mentions if snapshot.mentions > 0 else 0
            data['sentiment_sum'] += float(snapshot.sentiment_score) * snapshot.mentions if snapshot.mentions > 0 else 0
            data['visibility_sum'] += float(snapshot.visibility_score) * snapshot.mentions if snapshot.mentions > 0 else 0
            
            # Track platform-specific mentions for platform growth chart
            platform_key = platform_mapping.get(snapshot.platform, snapshot.platform.lower() if snapshot.platform else 'other')
            if platform_key not in data['platforms']:
                data['platforms'][platform_key] = 0
            data['platforms'][platform_key] += snapshot.mentions
        
        # Calculate weighted averages for each month
        for month_key, data in monthly_data.items():
            if data['total_mentions'] > 0:
                data['avg_position'] = round(data['position_sum'] / data['total_mentions'], 2)
                data['sentiment'] = round(data['sentiment_sum'] / data['total_mentions'], 2)
                data['visibility_score'] = round(data['visibility_sum'] / data['total_mentions'], 2)
            # Clean up helper fields
            del data['position_sum']
            del data['sentiment_sum']
            del data['visibility_sum']
            del data['total_mentions']
        
        # Build visibility trend (sorted by month)
        sorted_months = sorted(monthly_data.keys())
        visibility_trend = []
        for k in sorted_months:
            data = monthly_data[k]
            visibility_trend.append({
                'month': data['month'],
                'mentions': data['mentions'],
                'avg_position': data['avg_position'],
                'sentiment': data['sentiment'],
                'visibility_score': data['visibility_score']
            })
        
        # Platform growth (aggregate by month)
        platform_growth = []
        for k in sorted_months:
            data = monthly_data[k]
            platforms = data.get('platforms', {})
            platform_growth.append({
                'month': data['month'],
                'chatgpt': platforms.get('chatgpt', 0),
                'claude': platforms.get('claude', 0),
                'perplexity': platforms.get('perplexity', 0),
                'gemini': platforms.get('gemini', 0),
            })
        
        # Calculate summary metrics (first vs last period)
        if len(visibility_trend) >= 2:
            first = visibility_trend[0]
            last = visibility_trend[-1]
            
            visibility_growth = round(((last['visibility_score'] - first['visibility_score']) / first['visibility_score'] * 100) if first['visibility_score'] > 0 else 0, 1)
            mention_growth = round(((last['mentions'] - first['mentions']) / first['mentions'] * 100) if first['mentions'] > 0 else 0, 1)
            position_improvement = round(((first['avg_position'] - last['avg_position']) / first['avg_position'] * 100) if first['avg_position'] > 0 else 0, 1)
        else:
            visibility_growth = 0
            mention_growth = 0
            position_improvement = 0
        
        # Competitor comparison (from ShareOfVoiceAnalytics)
        from analytics.models import ShareOfVoiceAnalytics
        competitor_data = {}
        sov_records = ShareOfVoiceAnalytics.objects.filter(
            domain_id=domain_id,
            timestamp__gte=start_date,
            timestamp__lte=end_date
        ).order_by('timestamp')
        
        for record in sov_records:
            month_key = get_month_key(record.timestamp)
            if month_key not in competitor_data:
                competitor_data[month_key] = {}
            competitor_name = record.competitor.name if record.competitor else 'Your Brand'
            # Normalize competitor name for frontend
            comp_key = competitor_name.lower().replace(' ', '').replace('-', '')
            competitor_data[month_key][comp_key] = record.mention_count
        
        competitor_comparison = []
        for k in sorted_months:
            comp_data = {'month': monthly_data[k]['month']}
            if k in competitor_data:
                comp_data.update(competitor_data[k])
            else:
                comp_data['yourbrand'] = monthly_data[k]['mentions']
            competitor_comparison.append(comp_data)
        
        # Calculate market share gain
        market_share_gain = 0
        if len(competitor_comparison) >= 2:
            first_comp = competitor_comparison[0]
            last_comp = competitor_comparison[-1]
            first_total = sum(v for k, v in first_comp.items() if k != 'month')
            last_total = sum(v for k, v in last_comp.items() if k != 'month')
            first_share = (first_comp.get('yourbrand', 0) / first_total * 100) if first_total > 0 else 0
            last_share = (last_comp.get('yourbrand', 0) / last_total * 100) if last_total > 0 else 0
            market_share_gain = round(last_share - first_share, 1)
        
        # Seasonal patterns (calculate historical average by month name)
        # Get all historical data for the same months across different years
        # Query platform-specific snapshots and aggregate
        all_historical_snapshots = DomainMetricSnapshot.objects.filter(
            domain_id=domain_id,
            period_type=selected_period_type
        ).exclude(platform__isnull=True).exclude(platform='').order_by('snapshot_date')
        
        # Group by month name (Jan, Feb, etc.) across all years
        month_historical_avg = {}
        for snapshot in all_historical_snapshots:
            month_name = get_month_name(snapshot.snapshot_date)
            if month_name not in month_historical_avg:
                month_historical_avg[month_name] = {'total': 0, 'count': 0}
            month_historical_avg[month_name]['total'] += snapshot.mentions
            month_historical_avg[month_name]['count'] += 1
        
        # Calculate averages
        for month_name, data in month_historical_avg.items():
            month_historical_avg[month_name] = round(data['total'] / data['count'] if data['count'] > 0 else 0, 0)
        
        seasonal_pattern = []
        for k in sorted_months:
            data = monthly_data[k]
            month_name = data['month']
            avg_year = month_historical_avg.get(month_name, data['mentions'])
            seasonal_pattern.append({
                'month': month_name,
                'mentions': data['mentions'],
                'avgYear': int(avg_year)
            })
        
        # Performance Forecast - Calculate future projections
        forecast = []
        if len(visibility_trend) >= 3:  # Need at least 3 data points for forecasting
            # Extract mentions values for trend analysis
            mentions_values = [v['mentions'] for v in visibility_trend]
            
            # Calculate linear regression for mentions
            n = len(mentions_values)
            x = list(range(n))
            x_mean = sum(x) / n
            y_mean = sum(mentions_values) / n
            
            numerator = sum((x[i] - x_mean) * (mentions_values[i] - y_mean) for i in range(n))
            denominator = sum((x[i] - x_mean) ** 2 for i in range(n))
            
            if denominator != 0:
                slope_mentions = numerator / denominator
                intercept_mentions = y_mean - slope_mentions * x_mean
            else:
                slope_mentions = 0
                intercept_mentions = y_mean
            
            # Calculate standard deviation for confidence intervals
            if n > 1:
                std_dev_mentions = statistics.stdev(mentions_values) if len(mentions_values) > 1 else 0
            else:
                std_dev_mentions = 0
            
            # Generate forecast for next 3 months
            forecast_months = 3
            last_month_date = datetime.strptime(sorted_months[-1] + '-01', '%Y-%m-%d').date()
            
            # Include last actual data point
            last_actual = visibility_trend[-1]
            forecast.append({
                'month': last_actual['month'],
                'actual': last_actual['mentions'],
                'forecast': round(last_actual['mentions'], 0),
                'upper': round(last_actual['mentions'] + (1.645 * std_dev_mentions), 0),
                'lower': round(max(0, last_actual['mentions'] - (1.645 * std_dev_mentions)), 0)
            })
            
            # Generate future forecasts
            for i in range(1, forecast_months + 1):
                future_month_date = last_month_date + relativedelta(months=i)
                month_name = future_month_date.strftime('%b')
                future_x = n + i - 1
                
                # Forecast mentions
                forecast_mentions = slope_mentions * future_x + intercept_mentions
                forecast_mentions = max(0, forecast_mentions)  # Ensure non-negative
                
                # Calculate confidence intervals (90% = 1.645 standard deviations)
                upper_mentions = forecast_mentions + (1.645 * std_dev_mentions)
                lower_mentions = max(0, forecast_mentions - (1.645 * std_dev_mentions))
                
                forecast.append({
                    'month': month_name,
                    'actual': None,  # No actual data for future
                    'forecast': round(forecast_mentions, 0),
                    'upper': round(upper_mentions, 0),
                    'lower': round(lower_mentions, 0)
                })
        
        # Key Milestones - Detect significant events
        milestones = []
        if len(visibility_trend) >= 2:
            # Track record highs and significant achievements
            max_mentions = 0
            max_visibility = 0
            best_position = float('inf')
            milestones_list = []  # Use list instead of dict to allow multiple milestones per month
            
            for i, trend in enumerate(visibility_trend):
                month_key = sorted_months[i]
                mentions = trend['mentions']
                visibility = trend['visibility_score']
                position = trend['avg_position']
                
                # Record high mentions
                if mentions > max_mentions:
                    prev_max = max_mentions
                    max_mentions = mentions
                    # Check if it's a significant milestone (round numbers or large increases)
                    if mentions >= 100 and (mentions % 100 == 0 or (prev_max > 0 and mentions > prev_max * 1.5)):
                        milestones_list.append({
                            'date': month_key + '-01',
                            'title': f'Reached {int(mentions)} mentions',
                            'description': f'Monthly mentions reached {int(mentions)}, a new record high',
                            'type': 'achievement',
                            'metric': 'mentions',
                            'value': int(mentions)
                        })
                
                # Record high visibility score
                if visibility > max_visibility:
                    prev_max_vis = max_visibility
                    max_visibility = visibility
                    if visibility >= 50 and (prev_max_vis == 0 or visibility > prev_max_vis * 1.2):
                        milestones_list.append({
                            'date': month_key + '-01',
                            'title': f'Visibility score: {visibility:.1f}',
                            'description': f'Reached visibility score of {visibility:.1f}, highest to date',
                            'type': 'achievement',
                            'metric': 'visibility_score',
                            'value': round(visibility, 1)
                        })
                
                # Best position (lowest is better)
                if position > 0 and position < best_position:
                    prev_best = best_position
                    best_position = position
                    if position <= 10 and (prev_best == float('inf') or position < prev_best * 0.8):
                        milestones_list.append({
                            'date': month_key + '-01',
                            'title': f'Average position: {position:.1f}',
                            'description': f'Improved to average position {position:.1f}, best performance yet',
                            'type': 'achievement',
                            'metric': 'position',
                            'value': round(position, 1)
                        })
                
                # Significant growth milestones
                if i > 0:
                    prev_mentions = visibility_trend[i-1]['mentions']
                    if prev_mentions > 0:
                        growth_rate = ((mentions - prev_mentions) / prev_mentions) * 100
                        if growth_rate >= 50:  # 50% or more growth
                            milestones_list.append({
                                'date': month_key + '-01',
                                'title': f'{growth_rate:.0f}% growth',
                                'description': f'Month-over-month growth of {growth_rate:.0f}% in mentions',
                                'type': 'growth',
                                'metric': 'mentions',
                                'value': round(growth_rate, 1)
                            })
            
            # Sort by date (most recent first) and limit to top 10
            milestones = sorted(
                milestones_list,
                key=lambda x: x['date'],
                reverse=True
            )[:10]
        
        return Response({
            'visibility_trend': visibility_trend,
            'platform_growth': platform_growth,
            'competitor_comparison': competitor_comparison,
            'seasonal_pattern': seasonal_pattern,
            'forecast': forecast,
            'milestones': milestones,
            'summary': {
                'visibility_growth': visibility_growth,
                'mention_growth': mention_growth,
                'position_improvement': position_improvement,
                'market_share_gain': market_share_gain
            },
            'period_months': months,
            'start_date': start_date.isoformat(),
            'end_date': end_date.isoformat()
        })
        
    except Exception as e:
        import traceback
        logger.error(f"Error in get_historical_trends: {str(e)}\n{traceback.format_exc()}")
        return Response(
            {'error': f'Failed to retrieve historical trends: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['POST'])
@permission_classes([AllowAny])  # Temporarily allow all for testing
def generate_prompt_variants(request):
    """
    Generate prompt variants using AI based on a main prompt or prompt group
    """
    try:
        main_prompt = request.data.get('main_prompt') or request.data.get('prompt')
        group_id = request.data.get('group_id')
        
        # If group_id is provided, try to get the primary prompt from the group
        if group_id and not main_prompt:
            try:
                group = PromptGroup.objects.get(id=group_id)
                primary_prompt_obj = group.prompts.filter(type='primary').order_by('created_at').first()
                if primary_prompt_obj:
                    main_prompt = primary_prompt_obj.prompt
                else:
                    # If no primary prompt found, try to get any prompt from the group
                    any_prompt = group.prompts.order_by('created_at').first()
                    if any_prompt:
                        main_prompt = any_prompt.prompt
                    else:
                        return Response(
                            {'error': f'Prompt group {group_id} has no prompts. Please add a prompt first.'},
                            status=status.HTTP_400_BAD_REQUEST
                        )
            except PromptGroup.DoesNotExist:
                return Response(
                    {'error': f'Prompt group with id {group_id} not found'},
                    status=status.HTTP_404_NOT_FOUND
                )
        
        # Validate that we have a main_prompt
        if not main_prompt:
            return Response(
                {'error': 'main_prompt is required (or provide group_id to get it from the group)'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        # Generate variants using OpenAI
        try:
            openai_client = get_openai_client()
        except Exception as e:
            logger.error(f"Failed to initialize OpenAI client: {e}")
            return Response(
                {'error': f'OpenAI client not available: {str(e)}. Please configure OPENAI_API_KEY in settings.'},
                status=status.HTTP_503_SERVICE_UNAVAILABLE
            )
        
        # Create the prompt for AI
        system_prompt = """You are an AI that generates prompt variants for monitoring brand mentions in AI chatbots.

Your task is to generate 5-8 relevant prompt variants that are similar to the main prompt but capture different ways users might phrase the same query.

Guidelines:
- Variants should maintain the core intent but vary in phrasing, formality, and specificity
- Include variations with synonyms, different word orders, and related contexts
- Keep variants natural and realistic
- Each variant should be a complete, searchable query
- Return ONLY a JSON array of strings, no other text

Example format:
["variant 1", "variant 2", "variant 3", ...]"""

        user_prompt = f'Generate prompt variants for this main prompt: "{main_prompt}"\n\nReturn only a JSON array of variant strings, no explanations or markdown.'

        try:
            response = openai_client.chat.completions.create(
                model="gpt-4o",
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt}
                ],
                temperature=0.7,
                max_tokens=500,
                timeout=30
            )
            
            ai_response = response.choices[0].message.content.strip()
            
            # Parse the response - it should be a JSON array
            # Remove markdown code blocks if present
            if ai_response.startswith('```'):
                # Extract JSON from code block
                lines = ai_response.split('\n')
                ai_response = '\n'.join(lines[1:-1]) if len(lines) > 2 else ai_response
            
            # Try to parse as JSON
            try:
                variants = json.loads(ai_response)
                if not isinstance(variants, list):
                    # If it's a dict with variants key
                    if 'variants' in variants:
                        variants = variants['variants']
                    else:
                        # Try to extract array from text
                        import re
                        array_match = re.search(r'\[.*?\]', ai_response, re.DOTALL)
                        if array_match:
                            variants = json.loads(array_match.group())
                        else:
                            raise ValueError("Could not parse variants from response")
            except json.JSONDecodeError:
                # Fallback: try to extract variants from text
                import re
                # Look for quoted strings
                variants = re.findall(r'"([^"]+)"', ai_response)
                if not variants:
                    # Try single quotes
                    variants = re.findall(r"'([^']+)'", ai_response)
                if not variants:
                    # Try numbered list
                    variants = re.findall(r'\d+\.\s*(.+?)(?=\n|$)', ai_response)
            
            # Clean and validate variants
            variants = [v.strip() for v in variants if v.strip() and len(v.strip()) > 5]
            
            if not variants:
                # Fallback: generate simple variants
                words = main_prompt.split()
                variants = [
                    f"best {main_prompt}",
                    f"top {main_prompt}",
                    f"affordable {main_prompt}",
                    f"{main_prompt} reviews",
                    f"{main_prompt} guide",
                ]
            
            return Response({
                'variants': variants[:8],  # Limit to 8 variants
                'count': len(variants[:8])
            })
            
        except Exception as e:
            logger.error(f"Error calling OpenAI API: {e}", exc_info=True)
            # Fallback: generate simple variants
            words = main_prompt.split()
            fallback_variants = [
                f"best {main_prompt}",
                f"top {main_prompt}",
                f"affordable {main_prompt}",
                f"{main_prompt} reviews",
                f"{main_prompt} guide",
                f"{main_prompt} comparison",
            ]
            return Response({
                'variants': fallback_variants,
                'count': len(fallback_variants),
                'warning': 'AI generation failed, using fallback variants'
            })
            
    except Exception as e:
        logger.error(f"Error generating prompt variants: {e}", exc_info=True)
        return Response(
            {'error': f'Failed to generate variants: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )

