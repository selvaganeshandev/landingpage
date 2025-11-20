"""
Excel Report Generator using openpyxl
"""
from io import BytesIO
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


class ExcelReportGenerator:
    """Generate Excel reports using openpyxl"""

    def __init__(self, report_data, report_type):
        self.data = report_data
        self.report_type = report_type
        self.workbook = openpyxl.Workbook()
        self.workbook.remove(self.workbook.active)  # Remove default sheet

    def generate(self):
        """Generate the Excel file and return BytesIO buffer"""
        # Create summary sheet
        self._create_summary_sheet()

        # Create data sheets based on report type
        if self.report_type == 'Executive Dashboard':
            self._create_executive_sheets()
        elif self.report_type == 'Detailed Analytics':
            self._create_analytics_sheets()
        elif self.report_type == 'Competitor Focus':
            self._create_competitor_sheets()
        elif self.report_type == 'Content Strategy':
            self._create_content_sheets()

        # Save to buffer
        buffer = BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        return buffer

    def _create_summary_sheet(self):
        """Create summary information sheet"""
        ws = self.workbook.create_sheet("Summary")

        # Header style
        header_fill = PatternFill(start_color="4A5568", end_color="4A5568", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=14)

        # Add title
        ws['A1'] = self.report_type
        ws['A1'].font = Font(bold=True, size=16)
        ws.merge_cells('A1:B1')

        # Add metadata
        ws['A3'] = 'Domain:'
        ws['B3'] = self.data.get('domain_name', 'N/A')

        ws['A4'] = 'Period Start:'
        ws['B4'] = self.data['period']['start'].strftime('%Y-%m-%d')

        ws['A5'] = 'Period End:'
        ws['B5'] = self.data['period']['end'].strftime('%Y-%m-%d')

        ws['A6'] = 'Generated:'
        ws['B6'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

        # Style metadata
        for row in range(3, 7):
            ws[f'A{row}'].font = Font(bold=True)

        # Auto-size columns
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = max_length + 2

    def _create_executive_sheets(self):
        """Create sheets for Executive Dashboard"""
        # Key Metrics Sheet
        ws = self.workbook.create_sheet("Key Metrics")
        metrics = self.data.get('key_metrics', {})

        ws['A1'] = 'Metric'
        ws['B1'] = 'Value'

        metrics_data = [
            ('Total Prompts', metrics.get('total_prompts', 0)),
            ('Total Mentions', metrics.get('total_mentions', 0)),
            ('Mention Rate', f"{metrics.get('mention_rate', 0)}%"),
            ('Average Sentiment', f"{metrics.get('avg_sentiment', 0):.2f}"),
            ('Competitors Tracked', metrics.get('competitor_count', 0)),
        ]

        for idx, (metric, value) in enumerate(metrics_data, start=2):
            ws[f'A{idx}'] = metric
            ws[f'B{idx}'] = value

        self._style_table(ws, 1, len(metrics_data) + 1)

        # Top Prompts Sheet
        ws = self.workbook.create_sheet("Top Prompts")
        top_prompts = self.data.get('top_prompts', [])

        ws['A1'] = 'Prompt'
        ws['B1'] = 'Type'
        ws['C1'] = 'Mentions'

        for idx, prompt in enumerate(top_prompts, start=2):
            ws[f'A{idx}'] = prompt.get('text', '')
            ws[f'B{idx}'] = prompt.get('type', 'N/A')
            ws[f'C{idx}'] = prompt.get('mentions', 0)

        self._style_table(ws, 1, len(top_prompts) + 1)

    def _create_analytics_sheets(self):
        """Create sheets for Detailed Analytics"""
        # LLM Performance Sheet
        ws = self.workbook.create_sheet("LLM Performance")
        llm_perf = self.data.get('llm_performance', {})

        ws['A1'] = 'LLM Model'
        ws['B1'] = 'Total Queries'
        ws['C1'] = 'Mentions'
        ws['D1'] = 'Avg Sentiment'

        row = 2
        for llm_name, stats in llm_perf.items():
            ws[f'A{row}'] = llm_name.upper()
            ws[f'B{row}'] = stats.get('total', 0)
            ws[f'C{row}'] = stats.get('mentions', 0)
            ws[f'D{row}'] = f"{stats.get('avg_sentiment', 0):.2f}"
            row += 1

        self._style_table(ws, 1, row - 1)

        # Daily Stats Sheet
        ws = self.workbook.create_sheet("Daily Statistics")
        daily_stats = self.data.get('daily_stats', [])

        ws['A1'] = 'Date'
        ws['B1'] = 'Total Queries'
        ws['C1'] = 'Mentions'

        for idx, stat in enumerate(daily_stats, start=2):
            ws[f'A{idx}'] = stat.get('date', '')
            ws[f'B{idx}'] = stat.get('total', 0)
            ws[f'C{idx}'] = stat.get('mentions', 0)

        self._style_table(ws, 1, len(daily_stats) + 1)

    def _create_competitor_sheets(self):
        """Create sheets for Competitor Focus"""
        ws = self.workbook.create_sheet("Competitors")
        competitors = self.data.get('competitors', [])

        ws['A1'] = 'Our Domain Mentions'
        ws['B1'] = self.data.get('our_mentions', 0)

        ws['A3'] = 'Competitor'
        ws['B3'] = 'Website'
        ws['C3'] = 'Mentions'

        for idx, comp in enumerate(competitors, start=4):
            ws[f'A{idx}'] = comp.get('name', '')
            ws[f'B{idx}'] = comp.get('website', '')
            ws[f'C{idx}'] = comp.get('mentions', 0)

        self._style_table(ws, 3, len(competitors) + 3)

    def _create_content_sheets(self):
        """Create sheets for Content Strategy"""
        # Content Gaps Sheet
        ws = self.workbook.create_sheet("Content Gaps")
        gaps = self.data.get('content_gaps', [])

        ws['A1'] = 'Prompt'
        ws['B1'] = 'Type'

        for idx, gap in enumerate(gaps, start=2):
            ws[f'A{idx}'] = gap.get('text', '')
            ws[f'B{idx}'] = gap.get('type', 'N/A')

        self._style_table(ws, 1, len(gaps) + 1)

        # Optimization Opportunities Sheet
        ws = self.workbook.create_sheet("Optimization")
        opportunities = self.data.get('optimization_opportunities', [])

        ws['A1'] = 'Prompt'
        ws['B1'] = 'Current Mention Rate'

        for idx, opp in enumerate(opportunities, start=2):
            ws[f'A{idx}'] = opp.get('text', '')
            ws[f'B{idx}'] = f"{opp.get('mention_rate', 0)}%"

        self._style_table(ws, 1, len(opportunities) + 1)

    def _style_table(self, ws, start_row, end_row):
        """Apply consistent styling to a table"""
        header_fill = PatternFill(start_color="4A5568", end_color="4A5568", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        # Style header row
        for cell in ws[start_row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Auto-size columns
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = min(max_length + 2, 60)
