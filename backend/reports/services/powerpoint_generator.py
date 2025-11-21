"""
PowerPoint Report Generator using python-pptx
"""
from io import BytesIO
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


class PowerPointReportGenerator:
    """Generate PowerPoint presentations using python-pptx"""

    def __init__(self, report_data, report_type):
        self.data = report_data
        self.report_type = report_type
        self.prs = Presentation()
        self.prs.slide_width = Inches(10)
        self.prs.slide_height = Inches(7.5)

    def generate(self):
        """Generate the PowerPoint file and return BytesIO buffer"""
        # Add title slide
        self._create_title_slide()

        # Add content slides based on report type
        if self.report_type == 'Executive Dashboard':
            self._create_executive_slides()
        elif self.report_type == 'Detailed Analytics':
            self._create_analytics_slides()
        elif self.report_type == 'Competitor Focus':
            self._create_competitor_slides()
        elif self.report_type == 'Content Strategy':
            self._create_content_slides()

        # Save to buffer
        buffer = BytesIO()
        self.prs.save(buffer)
        buffer.seek(0)
        return buffer

    def _create_title_slide(self):
        """Create title slide"""
        title_slide_layout = self.prs.slide_layouts[0]
        slide = self.prs.slides.add_slide(title_slide_layout)

        title = slide.shapes.title
        subtitle = slide.placeholders[1]

        title.text = self.report_type

        subtitle_text = f"""
{self.data.get('domain_name', 'N/A')}
Period: {self.data['period']['start'].strftime('%Y-%m-%d')} to {self.data['period']['end'].strftime('%Y-%m-%d')}
Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
        """
        subtitle.text = subtitle_text.strip()

    def _create_executive_slides(self):
        """Create slides for Executive Dashboard"""
        # Key Metrics Slide
        bullet_slide_layout = self.prs.slide_layouts[1]
        slide = self.prs.slides.add_slide(bullet_slide_layout)

        title = slide.shapes.title
        title.text = "Key Performance Indicators"

        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        metrics = self.data.get('key_metrics', {})

        p = tf.paragraphs[0]
        p.text = f"Total Prompts: {metrics.get('total_prompts', 0)}"

        metrics_list = [
            f"Total Mentions: {metrics.get('total_mentions', 0)}",
            f"Mention Rate: {metrics.get('mention_rate', 0)}%",
            f"Average Sentiment: {metrics.get('avg_sentiment', 0):.2f}",
            f"Competitors Tracked: {metrics.get('competitor_count', 0)}",
        ]

        for metric in metrics_list:
            p = tf.add_paragraph()
            p.text = metric
            p.level = 0

        # Top Prompts Slide
        slide = self.prs.slides.add_slide(bullet_slide_layout)
        title = slide.shapes.title
        title.text = "Top Performing Prompts"

        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame

        top_prompts = self.data.get('top_prompts', [])
        if top_prompts:
            p = tf.paragraphs[0]
            first_prompt = top_prompts[0]
            p.text = f"{first_prompt.get('text', '')[:50]}... ({first_prompt.get('mentions', 0)} mentions)"

            for prompt in top_prompts[1:5]:
                p = tf.add_paragraph()
                p.text = f"{prompt.get('text', '')[:50]}... ({prompt.get('mentions', 0)} mentions)"
                p.level = 0
        else:
            p = tf.paragraphs[0]
            p.text = "No prompt data available"

    def _create_analytics_slides(self):
        """Create slides for Detailed Analytics"""
        bullet_slide_layout = self.prs.slide_layouts[1]

        # LLM Performance Slide
        slide = self.prs.slides.add_slide(bullet_slide_layout)
        title = slide.shapes.title
        title.text = "LLM Performance Overview"

        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame

        llm_perf = self.data.get('llm_performance', {})
        if llm_perf:
            first = True
            for llm_name, stats in llm_perf.items():
                if first:
                    p = tf.paragraphs[0]
                    first = False
                else:
                    p = tf.add_paragraph()

                p.text = f"{llm_name.upper()}: {stats.get('mentions', 0)} mentions / {stats.get('total', 0)} queries"
                p.level = 0

        # Daily Trends Slide
        slide = self.prs.slides.add_slide(bullet_slide_layout)
        title = slide.shapes.title
        title.text = "Recent Daily Performance"

        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame

        daily_stats = self.data.get('daily_stats', [])[:7]  # Last 7 days
        if daily_stats:
            p = tf.paragraphs[0]
            first_stat = daily_stats[0]
            p.text = f"{first_stat.get('date', '')}: {first_stat.get('mentions', 0)} mentions"

            for stat in daily_stats[1:]:
                p = tf.add_paragraph()
                p.text = f"{stat.get('date', '')}: {stat.get('mentions', 0)} mentions"
                p.level = 0

    def _create_competitor_slides(self):
        """Create slides for Competitor Focus"""
        bullet_slide_layout = self.prs.slide_layouts[1]

        slide = self.prs.slides.add_slide(bullet_slide_layout)
        title = slide.shapes.title
        title.text = "Competitive Landscape"

        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame

        p = tf.paragraphs[0]
        p.text = f"Our Domain Mentions: {self.data.get('our_mentions', 0)}"

        competitors = self.data.get('competitors', [])[:8]  # Top 8
        for comp in competitors:
            p = tf.add_paragraph()
            p.text = f"{comp.get('name', '')}: {comp.get('mentions', 0)} mentions"
            p.level = 0

    def _create_content_slides(self):
        """Create slides for Content Strategy"""
        bullet_slide_layout = self.prs.slide_layouts[1]

        # Content Gaps Slide
        slide = self.prs.slides.add_slide(bullet_slide_layout)
        title = slide.shapes.title
        title.text = "Content Gaps"

        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame

        gaps = self.data.get('content_gaps', [])[:8]  # Top 8
        if gaps:
            p = tf.paragraphs[0]
            p.text = gaps[0].get('text', '')[:70] + '...'

            for gap in gaps[1:]:
                p = tf.add_paragraph()
                p.text = gap.get('text', '')[:70] + '...'
                p.level = 0
        else:
            p = tf.paragraphs[0]
            p.text = "No content gaps identified"

        # Optimization Opportunities Slide
        slide = self.prs.slides.add_slide(bullet_slide_layout)
        title = slide.shapes.title
        title.text = "Optimization Opportunities"

        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame

        opportunities = self.data.get('optimization_opportunities', [])[:8]  # Top 8
        if opportunities:
            p = tf.paragraphs[0]
            first_opp = opportunities[0]
            p.text = f"{first_opp.get('text', '')[:60]}... ({first_opp.get('mention_rate', 0)}% rate)"

            for opp in opportunities[1:]:
                p = tf.add_paragraph()
                p.text = f"{opp.get('text', '')[:60]}... ({opp.get('mention_rate', 0)}% rate)"
                p.level = 0
        else:
            p = tf.paragraphs[0]
            p.text = "All prompts performing well"
